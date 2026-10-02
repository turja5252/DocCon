# Created by Tanzim Nasir
# Copyright (c) 2026 Tanzim Nasir.
# Built for Elite Integrity Services.
# Unauthorized use by other companies is prohibited.
"""Transmittal log: fill TRANSMITTAL, File Transmittal, next CT/ST/FT number.

Does not run the Excel macros. Replays what File Transmittal already does:
copy the working sheet to a tab named with the current number, clear the list,
increment n. Client, Shop, and Field each have their own Excel letter.
"""
from __future__ import annotations

import contextlib
import os
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from doccon.jobs import find_job_folder, job_folder_keys
from doccon.kinds import BOOK_DIR, CLIENT, FIELD, PREFIX, SHOP
from doccon.log_layout import CLIENT_LAYOUT, BookLayout, layout_for
from doccon.match import MatchedRow
from doccon.pack_state import PackExtra, letter_status
from doccon.pep import DOC_CONTROL_FROM, PepCover, email_line
from doccon.transmittal_books import adopt_transmittal_books

CLIENT_OUT = BOOK_DIR[CLIENT]
WORKING_SHEET = "TRANSMITTAL"
TEMPLATE_BOOK_NAMES = ("CT-202X-XXX.xlsm", "CT-20XX-XXX.xlsm", "CT-202X-XXX.xlsx")
PLACEHOLDER_JOBS = {"", "202X-XXX", "20XX-XXX", "202X-0XX"}

DOC_START_ROW = CLIENT_LAYOUT.doc_start_row
PAGE1_LAST_ROW = CLIENT_LAYOUT.page1_last_row
PAGE2_LAST_ROW = CLIENT_LAYOUT.page2_last_row
PAGE3_LAST_ROW = CLIENT_LAYOUT.page3_last_row
MAX_LINES = PAGE3_LAST_ROW - DOC_START_ROW + 1  # 80

CLIENT_STATUSES = CLIENT_LAYOUT.statuses


class LogError(ValueError):
    """Raised when the client book cannot be filled or filed."""


@dataclass(frozen=True)
class LogLine:
    document_no: str
    rev: str
    description: str
    status: str
    notes: str = ""


@dataclass(frozen=True)
class ClientBook:
    path: Path
    job_number: str
    next_number: int
    filed_tabs: tuple[str, ...]

    kind: str = CLIENT

    @property
    def cover_id(self) -> str:
        return f"{PREFIX[self.kind]}-{self.job_number}-{self.next_number}"


@dataclass(frozen=True)
class FileClientResult:
    book: Path
    cover_id: str
    sheet_name: str
    line_count: int
    pdf_path: Path | None


@dataclass(frozen=True)
class BookCover:
    """TO / CC / project already typed on TRANSMITTAL (Sarah's letter)."""

    to_line: str = ""
    cc_line: str = ""
    project_description: str = ""


def client_status_from_jira(purpose: str) -> str:
    token = (purpose or "").strip().casefold().replace("_", " ")
    mapping = {
        "approval": "APPROVAL",
        "info": "INFORMATION",
        "information": "INFORMATION",
        "planned": "REVIEW",
        "na": "INFORMATION",
        "n/a": "INFORMATION",
        "construction": "CONSTRUCTION",
        "review": "REVIEW",
        "as-built": "AS-BUILT",
        "as built": "AS-BUILT",
    }
    return mapping.get(token, "APPROVAL")


def lines_from_rows(rows: list[MatchedRow], kind: str = CLIENT) -> list[LogLine]:
    lines: list[LogLine] = []
    for row in rows:
        drawing = row.drawing
        rev = ""
        if row.pdf and row.pdf.rev:
            rev = str(row.pdf.rev).strip()
        elif drawing.outgoing_rev:
            rev = str(drawing.outgoing_rev).strip()
        lines.append(
            LogLine(
                document_no=drawing.drawing_id or drawing.key,
                rev=rev,
                description=drawing.title or "",
                status=line_status(drawing.purpose, kind),
            )
        )
    return lines


def lines_from_extras(items: list[PackExtra] | tuple[PackExtra, ...], kind: str = CLIENT) -> list[LogLine]:
    """Letter lines for PDFs that have no Jira issue. They are not an EDDI row."""
    lines: list[LogLine] = []
    for item in items:
        document_no = (item.document_no or "").strip() or Path(item.path).stem
        lines.append(
            LogLine(
                document_no=document_no,
                rev=(item.rev or "").strip(),
                description=(item.description or "").strip(),
                status=letter_status(kind, item.status),
            )
        )
    return lines


def pages_needed(line_count: int, layout: BookLayout | None = None) -> int:
    form = layout or CLIENT_LAYOUT
    try:
        return form.pages_needed(line_count)
    except ValueError as exc:
        raise LogError(str(exc)) from exc


def shop_status_from_jira(purpose: str, *, statuses: tuple[str, ...] = ()) -> str:
    token = (purpose or "").strip().casefold().replace("_", " ")
    allowed = statuses or layout_for(CLIENT).statuses
    mapping = {
        "approval": "CONSTRUCTION",
        "info": "INFORMATION",
        "information": "INFORMATION",
        "planned": "INFORMATION",
        "na": "PURCHASING ONLY",
        "n/a": "PURCHASING ONLY",
        "construction": "CONSTRUCTION",
        "purchasing": "PURCHASING ONLY",
        "purchasing only": "PURCHASING ONLY",
        "review": "INFORMATION",
        "as-built": "CONSTRUCTION",
        "as built": "CONSTRUCTION",
    }
    mapped = mapping.get(token, "CONSTRUCTION")
    if mapped not in allowed:
        if mapped == "PURCHASING ONLY" and "INFORMATION" in allowed:
            return "INFORMATION"
        return allowed[0]
    return mapped


def line_status(purpose: str, kind: str = CLIENT) -> str:
    if kind == CLIENT:
        return client_status_from_jira(purpose)
    return shop_status_from_jira(purpose, statuses=layout_for(kind).statuses)


def find_client_book(job_folder: Path, job_number: str) -> Path | None:
    return find_book(job_folder, job_number, CLIENT)


def _exact_named_book(folder: Path, prefix: str, job: str) -> Path | None:
    for suffix in (".xlsm", ".xlsx"):
        candidate = folder / f"{prefix}-{job}{suffix}"
        if candidate.is_file() and not candidate.name.startswith("~$"):
            return candidate
    return None


def _named_book(job_folder: Path, job_number: str, kind: str) -> Path | None:
    """Exact CT|ST|FT-{job}.xlsm, then the parent job (`2026-077-1` → `CT-2026-077`)."""
    out = Path(job_folder) / BOOK_DIR[kind]
    if not out.is_dir():
        return None
    prefix = PREFIX[kind]
    for key in job_folder_keys(job_number):
        found = _exact_named_book(out, prefix, key)
        if found is not None:
            return found
    return None


def resolved_log_job(path: Path, typed_job: str, stored_job: str) -> str:
    """Job number to keep on the letter.

    A child Jira Job Number may file the parent book. The cell stays `2026-077`
    and the cover stays `CT-2026-077-n`. A different job still refuses.
    """
    typed = typed_job.strip()
    stored = (stored_job or "").strip()
    folded = {key.casefold(): key for key in job_folder_keys(typed)}
    if stored and stored not in PLACEHOLDER_JOBS:
        match = folded.get(stored.casefold())
        if match is None:
            raise LogError(f"{path.name} is for job {stored}, not {typed}.")
        return match
    return typed


def find_book(job_folder: Path, job_number: str, kind: str = CLIENT) -> Path | None:
    named = _named_book(job_folder, job_number, kind)
    if named is not None:
        return named
    out = Path(job_folder) / BOOK_DIR[kind]
    if not out.is_dir():
        return None
    prefix = PREFIX[kind]
    if kind == CLIENT:
        for name in TEMPLATE_BOOK_NAMES:
            candidate = out / name
            if candidate.is_file():
                return candidate
    found: list[Path] = []
    try:
        children = list(out.iterdir())
    except OSError:
        return None
    token = prefix.casefold() + "-"
    for child in children:
        if child.name.startswith("~$"):
            continue
        if child.suffix.casefold() not in {".xlsm", ".xlsx"}:
            continue
        if child.name.casefold().startswith(token):
            found.append(child)
    if len(found) == 1:
        return found[0]
    return None


def template_client_book() -> Path | None:
    return template_book(CLIENT)


def template_book(kind: str = CLIENT) -> Path | None:
    """Blank CT/ST/FT book from the jobs template — logo and Add Page / File buttons."""
    prefix = PREFIX[kind]
    names = (f"{prefix}-202X-XXX.xlsm", f"{prefix}-20XX-XXX.xlsm", f"{prefix}-202X-XXX.xlsx")
    for token in ("202X-0XX", "202X-XXX"):
        folder = find_job_folder(token)
        if folder is None:
            continue
        for name in names:
            candidate = folder / BOOK_DIR[kind] / name
            if candidate.is_file():
                return candidate
    return None


def prepare_client_book(job_folder: Path, job_number: str) -> Path:
    return prepare_book(job_folder, job_number, CLIENT)


def prepare_book(job_folder: Path, job_number: str, kind: str = CLIENT) -> Path:
    """Point at {CT|ST|FT}-{job}.xlsm, or the parent-job book when that is the file.

    A template is renamed to the typed Job Number only when no named book exists.
    `CT-2026-077.xlsm` is not renamed to `CT-2026-077-1.xlsm`.
    """
    job = job_number.strip()
    if not job:
        raise LogError("Job Number is required.")
    named = _named_book(job_folder, job, kind)
    if named is not None:
        return named
    adopt_transmittal_books(job_folder, job)
    dest = Path(job_folder) / BOOK_DIR[kind] / f"{PREFIX[kind]}-{job}.xlsm"
    if dest.is_file():
        return dest
    found = find_book(job_folder, job, kind)
    if found is None:
        expected = " or ".join(f"{PREFIX[kind]}-{key}.xlsm" for key in job_folder_keys(job))
        raise LogError(
            f"No {kind} transmittal book in {BOOK_DIR[kind]}. Expected {expected}."
        )
    return found


def inspect_client_book(path: Path, job_number: str) -> ClientBook:
    return inspect_book(path, job_number, CLIENT)


def inspect_book(path: Path, job_number: str, kind: str = CLIENT) -> ClientBook:
    job = job_number.strip()
    layout = layout_for(kind)
    wb = load_workbook(path, data_only=False, keep_vba=path.suffix.casefold() == ".xlsm")
    try:
        if WORKING_SHEET not in wb.sheetnames:
            raise LogError(f"{path.name} has no {WORKING_SHEET} tab.")
        ws = wb[WORKING_SHEET]
        stored_job = str(ws[layout.job_cell].value or "").strip()
        log_job = resolved_log_job(path, job, stored_job)
        next_number = _as_int(ws[layout.number_cell].value, default=1)
        filed = tuple(name for name in wb.sheetnames if name != WORKING_SHEET and _is_filed_tab(name))
    finally:
        wb.close()
    return ClientBook(path=path, job_number=log_job, next_number=next_number, filed_tabs=filed, kind=kind)


def read_book_cover(path: Path, kind: str = CLIENT) -> BookCover:
    """TO, CC, and project already on TRANSMITTAL. Empty cells stay empty."""
    layout = layout_for(kind)
    # Three cells only. keep_vba parses the macro blob; on an online-only Dropbox
    # file that holds the window until the whole xlsm downloads.
    wb = load_workbook(path, data_only=False, keep_vba=False)
    try:
        if WORKING_SHEET not in wb.sheetnames:
            return BookCover()
        ws = wb[WORKING_SHEET]
        to_raw = _cover_cell(ws, layout.to_cell)
        cc_raw = _cover_cell(ws, layout.cc_cell)
        project = _cover_cell(ws, layout.project_cell)
        return BookCover(
            to_line=email_line(to_raw),
            cc_line=email_line(cc_raw),
            project_description=project,
        )
    finally:
        wb.close()


def book_cover_for_job(job_folder: Path | None, job_number: str, kind: str = CLIENT) -> BookCover:
    if job_folder is None or not job_number.strip():
        return BookCover()
    path = find_book(job_folder, job_number, kind)
    if path is None:
        return BookCover()
    try:
        return read_book_cover(path, kind)
    except (OSError, KeyError, ValueError, TypeError):
        return BookCover()


def pick_cover_fields(
    *,
    book: BookCover,
    pack_to: str = "",
    pack_cc: str = "",
    pack_project: str = "",
    pep: PepCover | None = None,
    kind: str = CLIENT,
    field_to: str = "",
    field_cc: str = "",
    shop_to: str = "",
    shop_cc: str = "",
) -> BookCover:
    """Prefer the letter, then the saved pack, then the default for this kind.

    Client default is the PEP. Shop and Field defaults are their Settings lists.
    """
    if kind == FIELD:
        pep_to = field_to
        pep_cc = field_cc
    elif kind == SHOP:
        pep_to = shop_to
        pep_cc = shop_cc
    elif pep is not None and kind == CLIENT:
        pep_to = pep.to_line
        pep_cc = pep.cc_line
    else:
        pep_to = ""
        pep_cc = ""
    pep_project = pep.project_description if pep is not None else ""
    return BookCover(
        to_line=_first_filled(book.to_line, pack_to, pep_to),
        cc_line=_first_filled(book.cc_line, pack_cc, pep_cc),
        project_description=_first_filled(book.project_description, pack_project, pep_project),
    )


def file_client_transmittal(
    path: Path,
    job_number: str,
    lines: list[LogLine],
    *,
    issued: date | None = None,
    expected_return: date | str | None = None,
    cover: PepCover | None = None,
    kind: str = CLIENT,
) -> FileClientResult:
    return _commit_transmittal(
        path,
        job_number,
        lines,
        issued=issued,
        expected_return=expected_return,
        cover=cover,
        archive=True,
        kind=kind,
    )


def fill_client_transmittal(
    path: Path,
    job_number: str,
    lines: list[LogLine],
    *,
    issued: date | None = None,
    expected_return: date | str | None = None,
    cover: PepCover | None = None,
    kind: str = CLIENT,
) -> FileClientResult:
    """Write TRANSMITTAL only. Leave Add Page / File Transmittal for Excel."""
    return _commit_transmittal(
        path,
        job_number,
        lines,
        issued=issued,
        expected_return=expected_return,
        cover=cover,
        archive=False,
        kind=kind,
    )


def _commit_transmittal(
    path: Path,
    job_number: str,
    lines: list[LogLine],
    *,
    issued: date | None = None,
    expected_return: date | str | None = None,
    cover: PepCover | None = None,
    archive: bool,
    kind: str = CLIENT,
) -> FileClientResult:
    if not lines:
        raise LogError("Select at least one drawing for this transmittal.")
    layout = layout_for(kind)
    pages = pages_needed(len(lines), layout)
    job = job_number.strip()
    keep_vba = path.suffix.casefold() == ".xlsm"
    wb = load_workbook(path, data_only=False, keep_vba=keep_vba)
    try:
        if WORKING_SHEET not in wb.sheetnames:
            raise LogError(f"{path.name} has no {WORKING_SHEET} tab.")
        ws = wb[WORKING_SHEET]
        stored_job = str(ws[layout.job_cell].value or "").strip()
        job = resolved_log_job(path, job, stored_job)
        number = _as_int(ws[layout.number_cell].value, default=1)
        sheet_name = str(number)
        if archive and sheet_name in wb.sheetnames:
            raise LogError(f"{path.name} already has a tab named {sheet_name}.")
        if path.suffix.casefold() == ".xlsm" and os.name == "nt":
            wb.close()
            from doccon.excel_pdf import file_client_with_excel

            file_client_with_excel(
                path,
                job=job,
                number=number,
                pages=pages,
                lines=lines,
                issued=issued or date.today(),
                expected_return=expected_return,
                cover=cover,
                template=template_book(kind),
                archive=archive,
                kind=kind,
            )
            cover_id = f"{PREFIX[kind]}-{job}-{number}"
            return FileClientResult(
                book=path,
                cover_id=cover_id,
                sheet_name=sheet_name if archive else WORKING_SHEET,
                line_count=len(lines),
                pdf_path=None,
            )
        _write_working_sheet(
            ws,
            layout=layout,
            job=job,
            number=number,
            pages=pages,
            lines=lines,
            issued=issued or date.today(),
            expected_return=expected_return,
            cover=cover,
        )
        if archive:
            archived = wb.copy_worksheet(ws)
            archived.title = sheet_name
            _strip_action_buttons(archived)
            current_index = wb.sheetnames.index(sheet_name)
            wb.move_sheet(archived, offset=1 - current_index)
            _reset_working_sheet(ws, layout=layout, next_number=number + 1)
        wb.save(path)
    finally:
        with contextlib.suppress(Exception):
            wb.close()
    cover_id = f"{PREFIX[kind]}-{job}-{number}"
    return FileClientResult(
        book=path,
        cover_id=cover_id,
        sheet_name=sheet_name if archive else WORKING_SHEET,
        line_count=len(lines),
        pdf_path=None,
    )


def cover_date_is_na(text: str) -> bool:
    """True for empty cover boxes and the N/A picker token (same meaning)."""
    token = (text or "").strip()
    return not token or token.casefold() == "n/a"


def parse_issued_date(text: str, *, default: date | None = None) -> date:
    token = (text or "").strip()
    if cover_date_is_na(token):
        return default or date.today()
    parsed = _try_date(token)
    if parsed is None:
        raise LogError(f"Date issued must be a date (for example {(default or date.today()):%Y-%m-%d}).")
    return parsed


def parse_expected_return(text: str) -> date | str:
    token = (text or "").strip()
    if cover_date_is_na(token):
        return "N/A"
    parsed = _try_date(token)
    return parsed if parsed is not None else token


def cover_issued_default(*, today: date | None = None) -> str:
    """Console Date issued on open / Load / Cancel Next. Always a calendar day, never N/A."""
    return (today or date.today()).isoformat()


def expected_return_from_issued(
    issued_text: str, days: int, *, today: date | None = None
) -> str:
    """Calendar-day offset from Date issued. Blank/N/A issued uses today. days=0 is same day."""
    base = parse_issued_date(issued_text, default=today or date.today())
    return (base + timedelta(days=int(days))).isoformat()


def cover_date_stamp(text: str) -> str:
    """YYYY-MM-DD when the cover box is a calendar day; else blank.

    Used for Expected return → Return Request Date and Date issued → Submission Date.
    Empty / N/A are the same: Pack does not stamp. Changing the cover box to N/A
    restores packed Next for that field to Now (see DrawingBoard.apply_cover_date_change).
    Free text stays on the letter only.
    """
    if cover_date_is_na(text):
        return ""
    parsed = _try_date((text or "").strip())
    return parsed.isoformat() if parsed else ""


def cover_return_request_stamp(text: str) -> str:
    """Expected return → packed Next Return Request Date. Blank does not wipe."""
    return cover_date_stamp(text)


def cover_submission_date_stamp(text: str) -> str:
    """Date issued → packed Next Submission Date. Blank does not wipe."""
    return cover_date_stamp(text)


def _try_date(text: str) -> date | None:
    token = text.strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d-%b-%Y", "%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(token, fmt).date()
        except ValueError:
            continue
    return None


def _date_cell(value: date) -> datetime:
    return datetime(value.year, value.month, value.day)


def _expected_cell(value: date | str | None) -> object:
    if value is None or value == "":
        return "N/A"
    if isinstance(value, date) and not isinstance(value, datetime):
        return _date_cell(value)
    if isinstance(value, datetime):
        return value
    return str(value)


def _write_working_sheet(
    ws: Worksheet,
    *,
    layout: BookLayout,
    job: str,
    number: int,
    pages: int,
    lines: list[LogLine],
    issued: date,
    expected_return: date | str | None = None,
    cover: PepCover | None = None,
) -> None:
    ws[layout.number_cell] = number
    ws[layout.pages_cell] = pages
    ws[layout.job_cell] = job
    if layout.from_cell:
        _set_value(ws, layout.from_cell, DOC_CONTROL_FROM)
    if cover is not None:
        if layout.to_cell:
            _set_value(ws, layout.to_cell, cover.to_line or None)
        if layout.cc_cell:
            _set_value(ws, layout.cc_cell, cover.cc_line or None)
        if layout.project_cell:
            _set_value(ws, layout.project_cell, cover.project_description or None)
        if layout.customer_cell:
            _set_value(ws, layout.customer_cell, cover.client or None)
        if layout.site_cell:
            _set_value(ws, layout.site_cell, cover.site or None)
    if layout.issued_cell:
        _set_value(ws, layout.issued_cell, _date_cell(issued))
    if layout.expected_cell:
        _set_value(ws, layout.expected_cell, _expected_cell(expected_return))
    last_row = (
        layout.page3_last_row if pages >= 3 else layout.page2_last_row if pages >= 2 else layout.page1_last_row
    )
    cols = (layout.doc_no_col, layout.rev_col, layout.desc_col, layout.status_col, layout.notes_col)
    for row in range(layout.doc_start_row, layout.page3_last_row + 1):
        ws.row_dimensions[row].hidden = row > last_row
        for col in cols:
            _set_value(ws, ws.cell(row, col).coordinate, None)
    for index, line in enumerate(lines):
        row = layout.doc_start_row + index
        _set_value(ws, ws.cell(row, layout.doc_no_col).coordinate, line.document_no)
        _set_value(ws, ws.cell(row, layout.rev_col).coordinate, _rev_value(line.rev))
        _set_value(ws, ws.cell(row, layout.desc_col).coordinate, line.description)
        _set_value(ws, ws.cell(row, layout.status_col).coordinate, line.status)
        _set_value(ws, ws.cell(row, layout.notes_col).coordinate, line.notes or None)


def _reset_working_sheet(ws: Worksheet, *, layout: BookLayout, next_number: int) -> None:
    ws[layout.number_cell] = next_number
    ws[layout.pages_cell] = 1
    for coord in layout.reset_cells:
        _set_value(ws, coord, None)
    cols = (layout.doc_no_col, layout.rev_col, layout.desc_col, layout.status_col, layout.notes_col)
    for row in range(layout.doc_start_row, layout.page3_last_row + 1):
        ws.row_dimensions[row].hidden = row > layout.page1_last_row
        for col in cols:
            _set_value(ws, ws.cell(row, col).coordinate, None)


def _set_value(ws: Worksheet, coord: str, value: object) -> None:
    cell = ws[coord]
    if isinstance(cell, MergedCell):
        return
    cell.value = value


def _cover_cell(ws: Worksheet, coord: str) -> str:
    if not coord:
        return ""
    cell = ws[coord]
    value = getattr(cell, "value", None)
    if value is None:
        return ""
    text = re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()
    if text.casefold() in {"", "n/a", "na", "tbd", "none", "nil"}:
        return ""
    return text


def _first_filled(*values: str) -> str:
    for value in values:
        token = (value or "").strip()
        if token:
            return token
    return ""


def _strip_action_buttons(ws: Worksheet) -> None:
    try:
        drawing = getattr(ws, "_drawing", None)
        anchors = getattr(drawing, "twoCellAnchor", None) if drawing is not None else None
        if not anchors:
            return
        keep = []
        for anchor in list(anchors):
            sp = getattr(anchor, "sp", None)
            macro = str(getattr(sp, "macro", "") or "")
            if "INSERT_PG" in macro.upper() or "counter" in macro.casefold():
                continue
            keep.append(anchor)
        if len(keep) != len(anchors):
            anchors[:] = keep
    except (AttributeError, TypeError, ValueError):
        return


def _as_int(value: object, default: int) -> int:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return default
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    text = str(value).strip()
    if re.fullmatch(r"\d+", text):
        return int(text)
    return default


def _rev_value(rev: str) -> str | int | None:
    token = (rev or "").strip()
    if not token:
        return None
    if re.fullmatch(r"\d+", token):
        return int(token)
    return token


def _is_filed_tab(name: str) -> bool:
    return bool(re.fullmatch(r"\d+", name.strip()))


def empty_client_workbook() -> Workbook:
    """Minimal TRANSMITTAL layout for tests (not the Elite branded book)."""
    wb = Workbook()
    ws = wb.active
    ws.title = WORKING_SHEET
    ws["E1"] = '="CT-"&A12&"-"&A2'
    ws["A2"] = 1
    ws["A3"] = 1
    ws["C3"] = "FROM:"
    ws["C4"] = DOC_CONTROL_FROM
    ws["C5"] = "TO:"
    ws["C7"] = "CC:"
    ws["A9"] = "PROJECT DESCRIPTON:"
    ws["A11"] = "ELITE JOB NUMBER:"
    ws["C11"] = "DATE ISSUED:"
    ws["D11"] = "EXPECTED DOCUMENT RETURNED DATE:"
    ws["A12"] = "202X-XXX"
    ws["A14"] = "DOCUMENT NO."
    ws["B14"] = "REV"
    ws["C14"] = "DESCRIPTION"
    ws["D14"] = "STATUS"
    ws["E14"] = "NOTES/COMMENTS"
    ws["A95"] = "GENERAL NOTES/COMMENTS"
    return wb


def empty_shop_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = WORKING_SHEET
    ws["K1"] = '="ST-"&I4&"-"&C2'
    ws["C2"] = 1
    ws["D2"] = 1
    ws["I3"] = "JOB NUMBER:"
    ws["J3"] = "CUSTOMER:"
    ws["I4"] = "202X-XXX"
    ws["A5"] = "ELITE CONTACT:"
    ws["I5"] = "DATE ISSUED:"
    ws["J5"] = "JOB SITE LOCATION:"
    ws["A7"] = "LIST OF ATTACHED DOCUMENT(S)"
    ws["A8"] = "DOCUMENT NO."
    ws["G8"] = "REV"
    ws["I8"] = "DESCRIPTION"
    ws["J8"] = "STATUS"
    ws["K8"] = "NOTES/COMMENTS"
    ws["A89"] = "NOTES / COMMENTS:"
    return wb


def empty_field_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = WORKING_SHEET
    ws["K1"] = '="FT-"&I6&"-"&C2'
    ws["C2"] = 1
    ws["D2"] = 1
    ws["I3"] = "TO:"
    ws["I5"] = "JOB NUMBER:"
    ws["J5"] = "CUSTOMER:"
    ws["I6"] = "202X-XXX"
    ws["A7"] = "ELITE CONTACT:"
    ws["I7"] = "DATE ISSUED:"
    ws["J7"] = "JOB SITE LOCATION:"
    ws["A9"] = "LIST OF ATTACHED DOCUMENT(S)"
    ws["A86"] = "NOTES / COMMENTS:"
    return wb
