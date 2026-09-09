# Created by Tanzim Nasir
# Copyright (c) 2026 Tanzim Nasir.
# Built for Elite Integrity Services.
# Unauthorized use by other companies is prohibited.
"""Rename blank CT/ST/FT template books to the real Job Number on load."""
from __future__ import annotations

import shutil
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from doccon.kinds import BOOK_DIR, CLIENT, FIELD, JOB_CELL, PREFIX, SHOP

WORKING_SHEET = "TRANSMITTAL"
PLACEHOLDER_JOBS = {"", "202X-XXX", "20XX-XXX", "202X-0XX"}
KINDS = (CLIENT, SHOP, FIELD)


@dataclass(frozen=True)
class AdoptedBook:
    kind: str
    path: Path
    renamed_from: str | None


def is_template_book_name(name: str, prefix: str) -> bool:
    """CT-202X-XXX.xlsm is a template. CT-2026-Tanzim.xlsm / CT-2026-075.xlsm are not."""
    stem = Path(name).stem
    token = prefix + "-"
    if not stem.upper().startswith(token.upper()):
        return False
    rest = stem[len(token) :]
    return "X" in rest.upper()


def adopt_transmittal_books(job_folder: Path, job_number: str) -> list[AdoptedBook]:
    """Rename CT-202X-XXX / ST-20XX-XXX / FT-20XX-XXX to CT|ST|FT-{job}.xlsm.

    Does not touch a book that already has the job name. Does not overwrite
    CT-2026-075.xlsm on a live job. Stamps the job number into a placeholder cell.
    """
    job = job_number.strip()
    if not job or not Path(job_folder).is_dir():
        return []
    adopted: list[AdoptedBook] = []
    for kind in KINDS:
        result = _adopt_one(Path(job_folder), job, kind)
        if result is not None:
            adopted.append(result)
    return adopted


def _adopt_one(job_folder: Path, job: str, kind: str) -> AdoptedBook | None:
    folder = job_folder / BOOK_DIR[kind]
    prefix = PREFIX[kind]
    dest = folder / f"{prefix}-{job}.xlsm"
    if dest.is_file():
        _stamp_job_number(dest, kind, job)
        return None
    source = _find_template(folder, prefix)
    if source is None:
        return None
    try:
        shutil.move(str(source), str(dest))
    except OSError:
        return None
    _stamp_job_number(dest, kind, job)
    return AdoptedBook(kind=kind, path=dest, renamed_from=source.name)


def _find_template(folder: Path, prefix: str) -> Path | None:
    if not folder.is_dir():
        return None
    try:
        children = list(folder.iterdir())
    except OSError:
        return None
    matches: list[Path] = []
    for child in children:
        if child.name.startswith("~$") or not child.is_file():
            continue
        if child.suffix.casefold() not in {".xlsm", ".xlsx"}:
            continue
        if is_template_book_name(child.name, prefix):
            matches.append(child)
    if len(matches) == 1:
        return matches[0]
    return None


def _has_vba(path: Path) -> bool:
    try:
        with ZipFile(path) as zipped:
            return "xl/vbaProject.bin" in zipped.namelist()
    except (OSError, BadZipFile):
        return False


def _stamp_job_number(path: Path, kind: str, job: str) -> None:
    # openpyxl save on a real .xlsm drops the logo and macro buttons.
    if _has_vba(path):
        return
    coord = JOB_CELL[kind]
    keep_vba = path.suffix.casefold() == ".xlsm"
    try:
        wb = load_workbook(path, data_only=False, keep_vba=keep_vba)
    except Exception:
        return
    try:
        if WORKING_SHEET not in wb.sheetnames:
            return
        ws = wb[WORKING_SHEET]
        cell = ws[coord]
        if isinstance(cell, MergedCell):
            return
        stored = str(cell.value or "").strip()
        if stored and stored not in PLACEHOLDER_JOBS and stored.casefold() != job.casefold():
            return
        if stored == job:
            return
        cell.value = job
        wb.save(path)
    except Exception:
        return
    finally:
        try:
            wb.close()
        except Exception:
            return
