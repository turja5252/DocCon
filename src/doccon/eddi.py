# Created by Tanzim Nasir
# Copyright (c) 2026 Tanzim Nasir.
# Built for Elite Integrity Services.
# Unauthorized use by other companies is prohibited.
"""Dated EDDI snapshot: fill a copy of the job's EDDI form (sheet Project).

Jira stays the register. The live ``EDDI-{job}.xlsm`` is the form and is not
overwritten. DocCon copies it to ``EDDI-{job}-{date}.xlsm``, fills groups 1–9
from console Next values (listed items with a paired PDF, not only Pack), and prints
``EDDI-{job}-{date}.pdf`` in ``3.0 Doc Con``. On that dated copy, printed
group titles are merged A:O and centered; empty 1–9 groups are hidden.
Submitted to Shop For is inserted before Shop Rev, and Submitted to Field For
before Field Rev, on the dated copy only.
Not attached to Outlook.
"""
from __future__ import annotations

import contextlib
import os
import re
import shutil
from copy import copy
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet

from doccon.client_log import LogError
from doccon.jobs import job_folder_keys
from doccon.match import MatchedRow
from doccon.register import DrawingRow, eddi_group_rank, eddi_group_title, is_generic_eddi, pack_sort_key

DOC_CON_DIR = Path("3.0 Doc Con")
PROJECT_SHEET = "Project"
ITEM_COLS = 15
PAPERSIZE_LETTER = 1
# Shop purpose is inserted at J, then Field purpose at M, on the dated copy.
PURPOSE_COLUMN_INSERTS = (10, 13)
CATEGORY_HEADER_MERGES = (
    "A3:A4",
    "B3:B4",
    "C3:C4",
    "D3:F3",
    "G3:I3",
    "J3:L3",
    "M3:O3",
)
PURPOSE_HEADERS = (
    (3, 10, "SHOP"),
    (4, 10, "SUBMITTED TO SHOP FOR"),
    (3, 13, "FIELD"),
    (4, 13, "SUBMITTED TO FIELD FOR"),
)
SHOP_PURPOSE_COL = 10
FIELD_PURPOSE_COL = 13
SHOP_PURPOSE_LIST = "IFC,IFI,IFU,Purchasing Only"
FIELD_PURPOSE_LIST = "IFC,IFI"
PURPOSE_COL_WIDTH = 14
_HEADER_RE = re.compile(r"^\s*(\d+)\.\s+")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATED_LIVE = re.compile(r"^EDDI-(.+)-(\d{4}-\d{2}-\d{2})\.(xlsx|xlsm)$", re.IGNORECASE)


@dataclass(frozen=True)
class EddiSnapshot:
    book: Path
    pdf: Path
    updated_rows: int
    warning: str = ""

    @property
    def note(self) -> str:
        base = f"EDDI: {self.pdf.name} ({self.updated_rows} item(s))"
        if self.warning:
            return f"{base}\n{self.warning}"
        return base


@dataclass
class FormGroup:
    number: int
    header_row: int
    junk_row: int | None
    first_item: int
    last_item: int

    @property
    def slot_count(self) -> int:
        if self.last_item < self.first_item:
            return 0
        return self.last_item - self.first_item + 1


@dataclass
class FillPlan:
    stamp: str
    inserts: list[tuple[int, int]] = field(default_factory=list)
    fills: list[tuple[int, list[str]]] = field(default_factory=list)
    hide: list[int] = field(default_factory=list)
    unhide: list[int] = field(default_factory=list)
    clear_rows: list[int] = field(default_factory=list)
    print_last_row: int = 1
    item_count: int = 0
    table_refs: list[tuple[str, str]] = field(default_factory=list)
    merge_headers: list[int] = field(default_factory=list)
    merge_ranges: list[str] = field(default_factory=list)
    drop_breaks: list[int] = field(default_factory=list)
    row_breaks: list[int] = field(default_factory=list)
    column_inserts: list[int] = field(default_factory=list)
    header_cells: list[tuple[int, int, str]] = field(default_factory=list)
    group_last_rows: list[int] = field(default_factory=list)
    shop_validation: str = ""
    field_validation: str = ""
    shop_purpose_col: int = 11
    field_purpose_col: int = 14


def row_has_eddi_pdf(row: MatchedRow) -> bool:
    """True when Locate… or auto-match paired a file that is still on disk."""
    hit = row.pdf
    if hit is None or (row.confidence or "").strip().casefold() == "missing":
        return False
    try:
        return Path(hit.path).is_file()
    except OSError:
        return False


def eddi_print_drawings(rows: list[MatchedRow]) -> list[DrawingRow]:
    """EDDI snapshot rows: paired PDF only. Caller still skips Generic / ungrouped."""
    return [row.drawing for row in rows if row_has_eddi_pdf(row)]


def grouped_eddi_items(drawings: list[DrawingRow]) -> list[tuple[str, list[DrawingRow]]]:
    """EDDI Status groups that have items. Generic is omitted. Empty groups are omitted."""
    visible = [row for row in drawings if not is_generic_eddi(row.eddi_status)]
    visible.sort(key=pack_sort_key)
    groups: list[tuple[str, list[DrawingRow]]] = []
    current = ""
    bucket: list[DrawingRow] = []
    for drawing in visible:
        title = eddi_group_title(drawing.eddi_status)
        if bucket and title != current:
            groups.append((current, bucket))
            bucket = []
        current = title
        bucket.append(drawing)
    if bucket:
        groups.append((current, bucket))
    return groups


def form_group_items(drawings: list[DrawingRow]) -> dict[int, list[DrawingRow]]:
    """Items for form groups 1–9. Generic and ungrouped are omitted."""
    buckets: dict[int, list[DrawingRow]] = {number: [] for number in range(1, 10)}
    visible = [row for row in drawings if not is_generic_eddi(row.eddi_status)]
    visible.sort(key=pack_sort_key)
    for drawing in visible:
        rank = eddi_group_rank(drawing.eddi_status)
        if 1 <= rank <= 9:
            buckets[rank].append(drawing)
    return buckets


def find_live_eddi_book(job_folder: Path, job_number: str) -> Path:
    """The job's EDDI form: ``EDDI-{job}.xlsm``, or the parent job's form.

    ``2026-077-1`` uses ``EDDI-2026-077-1.xlsm`` when that file exists, otherwise
    ``EDDI-2026-077.xlsm``. Dated snapshots are not the live form.
    """
    keys = job_folder_keys(job_number)
    if not keys:
        raise LogError("Job Number is required for the EDDI snapshot.")
    folder = Path(job_folder) / DOC_CON_DIR
    for key in keys:
        for suffix in (".xlsm", ".xlsx"):
            candidate = folder / f"EDDI-{key}{suffix}"
            if not candidate.is_file() or candidate.name.startswith("~$"):
                continue
            if is_dated_eddi_name(candidate.name, key):
                continue
            return candidate
    expected = " or ".join(f"EDDI-{key}.xlsm" for key in keys)
    raise LogError(
        f"No {expected} in 3.0 Doc Con. Put the EDDI form there (sheet Project), then try again."
    )


def parse_form_groups(ws: Worksheet) -> list[FormGroup]:
    last = _sheet_last_row(ws)
    headers: list[tuple[int, int]] = []
    for row in range(5, last + 1):
        number = _header_number(_cell_text(ws, row, 1))
        if number is not None:
            headers.append((number, row))
    groups: list[FormGroup] = []
    for index, (number, header_row) in enumerate(headers):
        block_end = headers[index + 1][1] - 1 if index + 1 < len(headers) else last
        junk: int | None = None
        first = header_row + 1
        if first <= block_end and _is_junk_row(ws, first):
            junk = first
            first = junk + 1
        last_item = block_end
        if first > last_item:
            last_item = first - 1
        groups.append(
            FormGroup(
                number=number,
                header_row=header_row,
                junk_row=junk,
                first_item=first,
                last_item=last_item,
            )
        )
    return groups


def build_fill_plan(ws: Worksheet, job_number: str, drawings: list[DrawingRow]) -> FillPlan:
    job = job_number.strip()
    groups = [group for group in parse_form_groups(ws) if 1 <= group.number <= 9]
    if not groups:
        raise LogError("The EDDI form has no Project groups 1–9.")
    buckets = form_group_items(drawings)
    tables = _table_bounds(ws)
    breaks = _row_break_ids(ws)
    inserts: list[tuple[int, int]] = []
    finalized: list[FormGroup] = []
    working = [
        FormGroup(g.number, g.header_row, g.junk_row, g.first_item, g.last_item) for g in groups
    ]
    for group in reversed(working):
        extra = max(0, len(buckets.get(group.number, [])) - group.slot_count)
        if extra:
            at = group.last_item + 1
            inserts.append((at, extra))
            group.last_item += extra
            _shift_below(finalized, at, extra)
            _shift_breaks(breaks, at, extra)
            for bounds in tables:
                _shift_table(bounds, at, extra)
        finalized.append(group)
    finalized.sort(key=lambda group: group.header_row)

    cols = _purpose_cols(ws)
    fills: list[tuple[int, list[str]]] = []
    hide: list[int] = []
    unhide: list[int] = []
    clear_rows: list[int] = []
    merge_headers: list[int] = []
    group_last_rows: list[int] = []
    count = 0
    last_row = 4
    for group in finalized:
        items = buckets.get(group.number, [])
        count += len(items)
        if group.junk_row is not None:
            hide.append(group.junk_row)
        if not items:
            hide.append(group.header_row)
            if group.slot_count:
                hide.extend(range(group.first_item, group.last_item + 1))
            continue
        unhide.append(group.header_row)
        merge_headers.append(group.header_row)
        if group.slot_count:
            clear_rows.extend(range(group.first_item, group.last_item + 1))
        for offset, drawing in enumerate(items):
            row = group.first_item + offset
            fills.append((row, item_values(drawing, cols)))
            unhide.append(row)
        group_last_rows.append(group.first_item + len(items) - 1)
        unused_start = group.first_item + len(items)
        if unused_start <= group.last_item:
            hide.extend(range(unused_start, group.last_item + 1))
        last_row = max(last_row, group.last_item, group.header_row)
    hide = _unique_keep(hide)
    unhide = _unique_keep(unhide)
    merge_headers = _unique_keep(merge_headers)
    hidden = set(hide) - set(unhide)
    keep_breaks = [row for row in breaks if not _is_stranded_break(row, hidden, last_row)]
    drop_breaks = [row for row in breaks if row not in set(keep_breaks)]
    add_columns = cols[1] == 0
    shop_letter = get_column_letter(cols[1]) if cols[1] else ""
    field_letter = get_column_letter(cols[4]) if cols[4] else ""
    filled_rows = sorted({row for row, _values in fills})
    if add_columns:
        for bounds in tables:
            if bounds[3] >= PURPOSE_COLUMN_INSERTS[0]:
                bounds[3] += len(PURPOSE_COLUMN_INSERTS)
    return FillPlan(
        stamp=f"EDDI-{job}",
        inserts=inserts,
        fills=fills,
        hide=hide,
        unhide=unhide,
        clear_rows=clear_rows,
        print_last_row=max(last_row, 1),
        item_count=count,
        table_refs=[
            (name, f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}")
            for name, min_col, min_row, max_col, max_row in tables
        ],
        merge_headers=merge_headers,
        merge_ranges=_snapshot_merges(add_columns),
        drop_breaks=drop_breaks,
        row_breaks=keep_breaks,
        column_inserts=list(PURPOSE_COLUMN_INSERTS) if add_columns else [],
        header_cells=list(PURPOSE_HEADERS) if add_columns else [],
        group_last_rows=group_last_rows,
        shop_validation=(
            ""
            if not shop_letter or _column_has_list(ws, shop_letter)
            else _retarget_validation(ws, shop_letter) or _ranges_for_rows(filled_rows, shop_letter)
        ),
        field_validation=(
            ""
            if not field_letter or _column_has_list(ws, field_letter)
            else _retarget_validation(ws, field_letter) or _ranges_for_rows(filled_rows, field_letter)
        ),
        shop_purpose_col=cols[1] or SHOP_PURPOSE_COL,
        field_purpose_col=cols[4] or FIELD_PURPOSE_COL,
    )


def fill_project_sheet(ws: Worksheet, job_number: str, drawings: list[DrawingRow]) -> int:
    """Fill sheet Project in-memory (tests / plain xlsx). Returns item count."""
    expand_purpose_columns(ws)
    plan = build_fill_plan(ws, job_number, drawings)
    apply_fill_plan(ws, plan)
    return plan.item_count


def apply_fill_plan(ws: Worksheet, plan: FillPlan) -> None:
    _set_a2(ws, plan.stamp)
    for at, count in plan.inserts:
        if count > 0:
            ws.insert_rows(at, count)
    for name, ref in plan.table_refs:
        table = ws.tables.get(name) if hasattr(ws, "tables") else None
        if table is not None:
            table.ref = ref
    for row in plan.clear_rows:
        _clear_item_row(ws, row)
    for row, values in plan.fills:
        for col, raw in enumerate(values, start=1):
            token = (raw or "").strip()
            if not token:
                continue
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell):
                continue
            cell.value = _typed_value(token)
    shown = set(plan.unhide)
    for row in plan.hide:
        if row not in shown:
            ws.row_dimensions[row].hidden = True
    for row in plan.unhide:
        ws.row_dimensions[row].hidden = False
    _layout_snapshot_sheet(ws, plan)
    _style_purpose_columns(ws, plan)
    ws.print_area = f"A1:{get_column_letter(ITEM_COLS)}{plan.print_last_row}"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def plan_payload(plan: FillPlan) -> dict:
    return {
        "sheet": PROJECT_SHEET,
        "stamp": plan.stamp,
        "inserts": [{"at": at, "count": count} for at, count in plan.inserts],
        "clear_rows": list(plan.clear_rows),
        "fills": [{"row": row, "values": values} for row, values in plan.fills],
        "hide": list(plan.hide),
        "unhide": list(plan.unhide),
        "print_last": plan.print_last_row,
        "tables": [{"name": name, "ref": ref} for name, ref in plan.table_refs],
        "merge_headers": list(plan.merge_headers),
        "merge_ranges": list(plan.merge_ranges),
        "drop_breaks": list(plan.drop_breaks),
        "column_inserts": list(plan.column_inserts),
        "header_cells": [
            {"row": row, "col": col, "value": text} for row, col, text in plan.header_cells
        ],
        "item_cols": ITEM_COLS,
        "group_last_rows": list(plan.group_last_rows),
        "shop_validation": plan.shop_validation,
        "field_validation": plan.field_validation,
        "shop_list": SHOP_PURPOSE_LIST,
        "field_list": FIELD_PURPOSE_LIST,
        "shop_purpose_col": plan.shop_purpose_col,
        "field_purpose_col": plan.field_purpose_col,
    }


def _purpose_cols(ws: Worksheet) -> tuple[int, int, int, int, int, int]:
    """Shop rev, shop purpose, shop date, field rev, field purpose, field date.

    The Tanzim template is Rev, Submitted to, Date. An older insert put Submitted to first.
    A 13-column form has no purpose columns.
    """
    labels = {col: _cell_text(ws, 4, col).casefold() for col in range(10, 16)}
    if "submitted to shop" in labels.get(11, ""):
        return (10, 11, 12, 13, 14, 15)
    if "submitted to shop" in labels.get(10, ""):
        return (11, 10, 12, 14, 13, 15)
    return (10, 0, 11, 12, 0, 13)


def item_values(drawing: DrawingRow, cols: tuple[int, int, int, int, int, int] | None = None) -> list[str]:
    layout = cols or (11, 10, 12, 14, 13, 15)
    width = 15 if layout[1] else 13
    values = [""] * width
    head = [
        (drawing.drawing_id or drawing.key).strip(),
        (drawing.client_document_number or "").strip(),
        (drawing.title or "").strip(),
        (drawing.outgoing_rev or "").strip(),
        (drawing.purpose or "").strip(),
        _short_date(drawing.submission_date),
        (drawing.incoming_rev or "").strip(),
        (drawing.approval or "").strip(),
        _short_date(drawing.return_date),
    ]
    for index, token in enumerate(head):
        values[index] = token
    placed = {
        layout[0]: (drawing.shop_ifc_rev or "").strip(),
        layout[1]: (drawing.shop_purpose or "").strip(),
        layout[2]: _short_date(drawing.shop_ifc_date),
        layout[3]: (drawing.field_ifc_rev or "").strip(),
        layout[4]: (drawing.field_purpose or "").strip(),
        layout[5]: _short_date(drawing.field_ifc_date),
    }
    for col, token in placed.items():
        if col:
            values[col - 1] = token
    return values


def _retarget_validation(ws: Worksheet, letter: str) -> str:
    """Same item rows as Submitted to Client For, on column J or M."""
    for dv in ws.data_validations.dataValidation:
        sqref = str(dv.sqref or "")
        if re.search(r"\bE\d+", sqref):
            return re.sub(r"E(?=\d)", letter, sqref)
    return ""


def _ranges_for_rows(rows: list[int], letter: str) -> str:
    if not rows:
        return ""
    parts: list[str] = []
    start = prev = rows[0]
    for row in rows[1:]:
        if row == prev + 1:
            prev = row
            continue
        parts.append(f"{letter}{start}:{letter}{prev}" if start != prev else f"{letter}{start}")
        start = prev = row
    parts.append(f"{letter}{start}:{letter}{prev}" if start != prev else f"{letter}{start}")
    return " ".join(parts)


def expand_purpose_columns(ws: Worksheet) -> None:
    """Insert Shop and Field purpose columns without redrawing the rest of the grid.

    Only merges that cross the insert are opened. New cells copy the border of the
    rev column beside them, so item rows stay vertical lines, not a new box.
    """
    if not _purpose_columns_missing(ws):
        return
    for col in PURPOSE_COLUMN_INSERTS:
        for merged in list(ws.merged_cells.ranges):
            if merged.min_col <= col <= merged.max_col:
                ws.unmerge_cells(str(merged))
        ws.insert_cols(col)
    last_row = int(ws.max_row or 1)
    for row in range(5, last_row + 1):
        _copy_border(ws.cell(row, SHOP_PURPOSE_COL + 1), ws.cell(row, SHOP_PURPOSE_COL))
        _copy_border(ws.cell(row, FIELD_PURPOSE_COL + 1), ws.cell(row, FIELD_PURPOSE_COL))
    for row, col, text in PURPOSE_HEADERS:
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        cell.value = text
    for col in (SHOP_PURPOSE_COL, FIELD_PURPOSE_COL):
        head = ws.cell(3, col)
        if isinstance(head, MergedCell):
            continue
        head.font = Font(name="Calibri", size=11, bold=True)
        head.alignment = Alignment(horizontal="center", vertical="center")
    for col in (SHOP_PURPOSE_COL + 1, FIELD_PURPOSE_COL + 1):
        leftover = ws.cell(3, col)
        if isinstance(leftover, MergedCell):
            continue
        token = str(leftover.value or "").casefold()
        if token.startswith("shop") or token.startswith("field"):
            leftover.value = None
    ws.column_dimensions["J"].width = PURPOSE_COL_WIDTH
    ws.column_dimensions["M"].width = PURPOSE_COL_WIDTH
    last = get_column_letter(ITEM_COLS)
    for ref in (f"A1:{last}1", f"A2:{last}2", *CATEGORY_HEADER_MERGES):
        _ensure_merge(ws, ref)


def _copy_border(src, dest) -> None:
    border = src.border
    if any(side is not None and side.style for side in (border.left, border.right, border.top, border.bottom)):
        dest.border = Border(left=border.left, right=border.right, top=border.top, bottom=border.bottom)
        return
    thin = Side(style="thin")
    dest.border = Border(left=thin, right=thin)


def _style_purpose_columns(ws: Worksheet, plan: FillPlan) -> None:
    """Wrap a long purpose such as Purchasing Only, and close the last visible row."""
    thin = Side(style="thin")
    wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    watched = (
        (plan.shop_purpose_col, plan.shop_purpose_col),
        (plan.field_purpose_col, plan.field_purpose_col),
    )
    for row, values in plan.fills:
        for col, _same in watched:
            raw = values[col - 1] if col and len(values) >= col else ""
            text = (raw or "").strip()
            if len(text) <= 12:
                continue
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell):
                continue
            cell.alignment = wrap
            current = ws.row_dimensions[row].height or 15
            ws.row_dimensions[row].height = max(float(current), 30)
    for row in plan.group_last_rows:
        for col in range(1, ITEM_COLS + 1):
            cell = ws.cell(row, col)
            if isinstance(cell, MergedCell):
                continue
            border = cell.border
            cell.border = Border(left=border.left, right=border.right, top=border.top, bottom=thin)
    _add_list_validation(ws, plan.shop_validation, SHOP_PURPOSE_LIST)
    _add_list_validation(ws, plan.field_validation, FIELD_PURPOSE_LIST)


def _add_list_validation(ws: Worksheet, sqref: str, formula: str) -> None:
    areas = [part for part in (sqref or "").replace(",", " ").split() if part]
    if not areas:
        return
    dv = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True)
    for area in areas:
        dv.add(area)
    ws.add_data_validation(dv)


def _column_has_list(ws: Worksheet, letter: str) -> bool:
    return any(re.search(rf"\b{letter}\d+", str(dv.sqref or "")) for dv in ws.data_validations.dataValidation)


def _purpose_columns_missing(ws: Worksheet) -> bool:
    """True when row 4 has no Submitted to Shop For column yet."""
    return _purpose_cols(ws)[1] == 0


def _snapshot_merges(add_columns: bool) -> list[str]:
    """Category headers, plus the title rows once the dated copy grows to column O."""
    merges = list(CATEGORY_HEADER_MERGES)
    if not add_columns:
        return merges
    last = get_column_letter(ITEM_COLS)
    return [f"A1:{last}1", f"A2:{last}2", *merges]


def snapshot_eddi(
    job_folder: Path,
    job_number: str,
    drawings: list[DrawingRow],
    issued: date,
    *,
    print_pdf: bool = True,
) -> EddiSnapshot:
    """Copy the live EDDI form, fill Project groups 1–9, print a dated PDF."""
    job = job_number.strip()
    if not job:
        raise LogError("Job Number is required for the EDDI snapshot.")
    live = find_live_eddi_book(job_folder, job)
    book_job = live.stem[5:] if live.stem.upper().startswith("EDDI-") else job
    folder = Path(job_folder) / DOC_CON_DIR
    folder.mkdir(parents=True, exist_ok=True)
    stem = f"EDDI-{book_job}-{issued.isoformat()}"
    dest_book = folder / f"{stem}{live.suffix}"
    dest_pdf = folder / f"{stem}.pdf"
    if print_pdf:
        return _snapshot_with_excel(live, dest_book, dest_pdf, book_job, drawings)
    shutil.copy2(live, dest_book)
    wb = _open_book(dest_book)
    try:
        count = fill_project_sheet(_project_sheet(wb), book_job, drawings)
        wb.save(dest_book)
    except OSError as exc:
        raise LogError(f"Could not write {dest_book.name}. Close it if it is open in Excel.\n{exc}") from exc
    finally:
        wb.close()
    return EddiSnapshot(book=dest_book, pdf=dest_pdf, updated_rows=count)


def is_dated_eddi_name(name: str, job_number: str) -> bool:
    match = _DATED_LIVE.match(Path(name).name)
    if not match:
        return False
    return match.group(1).casefold() == job_number.strip().casefold()


def _snapshot_with_excel(
    live: Path,
    dest_book: Path,
    dest_pdf: Path,
    job: str,
    drawings: list[DrawingRow],
) -> EddiSnapshot:
    if os.name != "nt":
        raise LogError("EDDI snapshot needs Excel on Windows.")
    from doccon.diag import describe_path, log
    from doccon.excel_pdf import ExcelPrintError, doccon_temp_dir, fill_eddi_form_and_export_pdf

    tmp_dir = doccon_temp_dir("eddi-form-")
    local_book = tmp_dir / dest_book.name
    local_pdf = tmp_dir / dest_pdf.name
    try:
        log("INFO", "eddi", f"copy {describe_path(live)} -> {describe_path(local_book)}")
        shutil.copy2(live, local_book)
        _unblock(local_book)
        prepared = load_workbook(local_book, data_only=False, keep_vba=local_book.suffix.casefold() == ".xlsm")
        try:
            expand_purpose_columns(_project_sheet(prepared))
            prepared.save(local_book)
        finally:
            prepared.close()
        wb = load_workbook(local_book, data_only=False, keep_vba=False)
        try:
            plan = build_fill_plan(_project_sheet(wb), job, drawings)
        finally:
            wb.close()
        printed: Path | None = None
        print_error = ""
        try:
            printed = fill_eddi_form_and_export_pdf(local_book, local_pdf, plan_payload(plan))
        except ExcelPrintError as exc:
            print_error = str(exc)
        except LogError as exc:
            if local_book.is_file() and _looks_like_print_error(exc):
                book = _copy_out(local_book, dest_book)
                pdf, warn = _copy_pdf_out(local_pdf, dest_pdf)
                warning = warn or str(exc)
                return EddiSnapshot(
                    book=book, pdf=pdf, updated_rows=plan.item_count, warning=warning
                )
            raise
        book = _copy_out(local_book, dest_book)
        source_pdf = printed if printed and printed.is_file() else local_pdf
        pdf, warn = _copy_pdf_out(source_pdf, dest_pdf)
        warning = print_error or warn
        if warning and not pdf.is_file():
            warning = warning or (
                f"Wrote {book.name}, but could not print the EDDI PDF. "
                "Close that PDF if it is open in a reader, then try Create EDDI again."
            )
        return EddiSnapshot(
            book=book, pdf=pdf, updated_rows=plan.item_count, warning=warning
        )
    except LogError:
        raise
    except OSError as exc:
        raise LogError("Could not write the EDDI snapshot.") from exc
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _open_book(path: Path) -> Workbook:
    keep_vba = path.suffix.casefold() == ".xlsm"
    try:
        return load_workbook(path, data_only=False, keep_vba=keep_vba)
    except Exception as exc:
        raise LogError(
            f"Could not open {path.name} as an Excel workbook. "
            "Close it if it is open, or pick a desktop Excel file (not a PDF)."
        ) from exc


def _project_sheet(wb: Workbook) -> Worksheet:
    if PROJECT_SHEET in wb.sheetnames:
        return wb[PROJECT_SHEET]
    raise LogError(f"The EDDI form has no {PROJECT_SHEET} sheet.")


def _sheet_last_row(ws: Worksheet) -> int:
    printed = _print_area_last_row(getattr(ws, "print_area", None))
    used = int(ws.max_row or 1)
    return max(printed, used, 5)


def _print_area_last_row(area: object) -> int:
    text = str(area or "").replace("'", "")
    if "!" in text:
        text = text.split("!", 1)[1]
    text = text.replace("$", "")
    match = re.search(r":[A-Za-z]+(\d+)\s*$", text)
    return int(match.group(1)) if match else 0


def _header_number(value: str) -> int | None:
    match = _HEADER_RE.match(value or "")
    if not match:
        return None
    number = int(match.group(1))
    if 1 <= number <= 9:
        return number
    return None


def _is_junk_row(ws: Worksheet, row: int) -> bool:
    token = _cell_text(ws, row, 1).casefold()
    return token.startswith("column")


def _cell_text(ws: Worksheet, row: int, col: int) -> str:
    value = ws.cell(row, col).value
    if value is None:
        return ""
    return str(value).strip()


def _set_a2(ws: Worksheet, stamp: str) -> None:
    cell = ws.cell(2, 1)
    if not isinstance(cell, MergedCell):
        cell.value = stamp


def _clear_item_row(ws: Worksheet, row: int) -> None:
    for col in range(1, ITEM_COLS + 1):
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        if cell.value is not None:
            cell.value = None


def _typed_value(token: str) -> date | str:
    if _ISO_DATE.match(token):
        try:
            return date.fromisoformat(token)
        except ValueError:
            return token
    return token


def _short_date(value: str) -> str:
    return (value or "").strip()[:10]


def _table_bounds(ws: Worksheet) -> list[list]:
    found: list[list] = []
    tables = getattr(ws, "tables", None)
    if not tables:
        return found
    for name in tables:
        table = tables[name]
        ref = getattr(table, "ref", None)
        if not ref:
            continue
        min_col, min_row, max_col, max_row = range_boundaries(str(ref))
        found.append([str(name), int(min_col), int(min_row), int(max_col), int(max_row)])
    return found


def _shift_table(bounds: list, at: int, count: int) -> None:
    _name, _min_col, min_row, _max_col, max_row = bounds
    if min_row >= at:
        bounds[2] = min_row + count
        bounds[4] = max_row + count
    elif max_row >= at:
        bounds[4] = max_row + count


def _shift_below(groups: list[FormGroup], at: int, count: int) -> None:
    for group in groups:
        if group.header_row >= at:
            group.header_row += count
        if group.junk_row is not None and group.junk_row >= at:
            group.junk_row += count
        if group.first_item >= at:
            group.first_item += count
        if group.last_item >= at:
            group.last_item += count


def _unique_keep(rows: list[int]) -> list[int]:
    seen: set[int] = set()
    out: list[int] = []
    for row in rows:
        if row in seen:
            continue
        seen.add(row)
        out.append(row)
    return out


def _row_break_ids(ws: Worksheet) -> list[int]:
    breaks = getattr(ws, "row_breaks", None)
    items = getattr(breaks, "brk", None) if breaks is not None else None
    if not items:
        return []
    found: list[int] = []
    for brk in items:
        row = int(getattr(brk, "id", 0) or 0)
        if row > 0:
            found.append(row)
    return found


def _shift_breaks(breaks: list[int], at: int, count: int) -> None:
    for index, row in enumerate(breaks):
        if row >= at:
            breaks[index] = row + count


def _is_stranded_break(after_row: int, hidden: set[int], print_last: int) -> bool:
    if after_row < 1 or after_row >= print_last:
        return True
    before = any(row not in hidden for row in range(5, after_row + 1))
    after = any(row not in hidden for row in range(after_row + 1, print_last + 1))
    return (not before) or (not after)


def _layout_snapshot_sheet(ws: Worksheet, plan: FillPlan) -> None:
    for row in plan.merge_headers:
        _merge_group_header(ws, row)
    for ref in plan.merge_ranges:
        min_col, min_row, _max_col, _max_row = range_boundaries(ref)
        _ensure_merge(ws, ref)
        _center_cell(ws.cell(min_row, min_col))
    _apply_row_breaks(ws, plan.row_breaks)


def _merge_group_header(ws: Worksheet, row: int) -> None:
    fill = _row_fill(ws, row)
    ref = f"A{row}:{get_column_letter(ITEM_COLS)}{row}"
    _ensure_merge(ws, ref)
    cell = ws.cell(row, 1)
    _center_cell(cell)
    if fill is not None:
        cell.fill = fill


def _ensure_merge(ws: Worksheet, ref: str) -> None:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    target = (min_col, min_row, max_col, max_row)
    for existing in list(ws.merged_cells.ranges):
        box = (existing.min_col, existing.min_row, existing.max_col, existing.max_row)
        if box == target:
            return
        if existing.max_row < min_row or existing.min_row > max_row:
            continue
        if existing.max_col < min_col or existing.min_col > max_col:
            continue
        ws.unmerge_cells(str(existing))
    ws.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def _center_cell(cell) -> None:
    prev = cell.alignment
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=bool(getattr(prev, "wrap_text", False)),
        shrink_to_fit=bool(getattr(prev, "shrink_to_fit", False)),
        textRotation=getattr(prev, "textRotation", 0) or 0,
        indent=getattr(prev, "indent", 0) or 0,
    )


def _row_fill(ws: Worksheet, row: int):
    for col in range(1, ITEM_COLS + 1):
        fill = ws.cell(row, col).fill
        if fill is not None and fill.fill_type not in (None, "none"):
            return copy(fill)
    return None


def _apply_row_breaks(ws: Worksheet, keep: list[int]) -> None:
    ws.row_breaks.brk = []
    for row in keep:
        if row > 0:
            ws.row_breaks.append(Break(id=int(row), man=True, max=ITEM_COLS))


def _looks_like_print_error(exc: BaseException) -> bool:
    low = str(exc).casefold()
    return "could not print" in low or "print the pdf" in low or "print the eddi pdf" in low


def _copy_out(source: Path, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    try:
        if dest.exists():
            dest.unlink()
        shutil.copy2(source, dest)
        return dest.resolve()
    except OSError as exc:
        from doccon.excel_pdf import _is_sharing_violation, _unique_sibling

        if not _is_sharing_violation(exc):
            raise
        alt = _unique_sibling(dest)
        shutil.copy2(source, alt)
        return alt.resolve()


def _copy_pdf_out(source: Path, dest: Path) -> tuple[Path, str]:
    if not source.is_file() or source.stat().st_size <= 0:
        return dest, (
            "Wrote the EDDI workbook, but could not print the PDF. "
            "Close that PDF if it is open in a reader, then try Create EDDI again."
        )
    try:
        written = _copy_out(source, dest)
    except OSError:
        return dest, (
            f"Wrote the EDDI workbook, but could not save {dest.name}. "
            "Close that PDF if it is open in a reader, then try Create EDDI again."
        )
    if written.resolve() != dest.resolve():
        return written, (
            f"{dest.name} was already open, so DocCon saved {written.name} instead. "
            "Close the old PDF in your reader."
        )
    return written, ""


def _unblock(path: Path) -> None:
    with contextlib.suppress(OSError):
        os.remove(f"{path}:Zone.Identifier")
