# Created by Tanzim Nasir
# Copyright (c) 2026 Tanzim Nasir.
# Built for Elite Integrity Services.
# Unauthorized use by other companies is prohibited.
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from doccon.client_log import (
    CLIENT_OUT,
    BookCover,
    LogError,
    LogLine,
    client_status_from_jira,
    cover_date_stamp,
    cover_issued_default,
    cover_return_request_stamp,
    cover_submission_date_stamp,
    empty_client_workbook,
    empty_field_workbook,
    empty_shop_workbook,
    expected_return_from_issued,
    file_client_transmittal,
    fill_client_transmittal,
    find_client_book,
    inspect_book,
    inspect_client_book,
    line_status,
    lines_from_rows,
    pages_needed,
    parse_expected_return,
    parse_issued_date,
    pick_cover_fields,
    prepare_client_book,
    read_book_cover,
)
from doccon.kinds import FIELD, SHOP
from doccon.log_layout import FIELD_LAYOUT, SHOP_LAYOUT
from doccon.match import MatchedRow, PdfHit
from doccon.pep import DOC_CONTROL_FROM, PepCover
from doccon.register import DrawingRow


def _drawing(**kwargs) -> DrawingRow:
    values = {
        "key": "P2024-1",
        "summary": "2026-Tanzim-1-1 Drawing-1",
        "drawing_id": "2026-Tanzim-1-1",
        "title": "Drawing-1",
        "status": "To Do",
        "job_number": "2026-Tanzim",
        "outgoing_rev": "A",
        "purpose": "Approval",
        "parent_summary": "Drawing Package",
    }
    values.update(kwargs)
    return DrawingRow(**values)


def test_status_maps_jira_purpose() -> None:
    assert client_status_from_jira("Approval") == "APPROVAL"
    assert client_status_from_jira("Info") == "INFORMATION"
    assert client_status_from_jira("") == "APPROVAL"
    assert line_status("Info", SHOP) == "INFORMATION"
    assert line_status("Approval", SHOP) == "CONSTRUCTION"
    assert line_status("NA", SHOP) == "PURCHASING ONLY"
    assert line_status("NA", FIELD) == "INFORMATION"


def test_pages_needed() -> None:
    assert pages_needed(1) == 1
    assert pages_needed(18) == 1
    assert pages_needed(19) == 2
    assert pages_needed(80) == 3
    assert pages_needed(20, SHOP_LAYOUT) == 1
    assert pages_needed(25, SHOP_LAYOUT) == 2
    assert pages_needed(23, FIELD_LAYOUT) == 1


def test_lines_prefer_pdf_rev() -> None:
    drawing = _drawing(outgoing_rev="A")
    hit = PdfHit(path=Path("2026-Tanzim-1-1 REV 0.pdf"), drawing_id="2026-Tanzim-1-1", rev="0")
    lines = lines_from_rows([MatchedRow(drawing=drawing, pdf=hit, confidence="High")])
    assert lines[0].rev == "0"
    assert lines[0].status == "APPROVAL"


def test_lines_blank_description_when_title_empty() -> None:
    drawing = _drawing(summary="2026-075-ITP-1-1", drawing_id="2026-075-ITP-1-1", title="")
    lines = lines_from_rows([MatchedRow(drawing=drawing, pdf=None, confidence="Missing")])
    assert lines[0].document_no == "2026-075-ITP-1-1"
    assert lines[0].description == ""


def test_find_prefers_named_book(tmp_path: Path) -> None:
    out = tmp_path / CLIENT_OUT
    out.mkdir(parents=True)
    (out / "CT-202X-XXX.xlsm").write_bytes(b"old")
    wanted = out / "CT-2026-Tanzim.xlsm"
    wanted.write_bytes(b"new")
    assert find_client_book(tmp_path, "2026-Tanzim") == wanted


def test_find_child_job_uses_parent_book(tmp_path: Path) -> None:
    out = tmp_path / CLIENT_OUT
    out.mkdir(parents=True)
    parent = out / "CT-2026-077.xlsm"
    parent.write_bytes(b"live")
    (out / "CT-2026-077-notes.xlsm").write_bytes(b"extra")
    assert find_client_book(tmp_path, "2026-077-1") == parent


def test_find_child_book_wins_over_parent(tmp_path: Path) -> None:
    out = tmp_path / CLIENT_OUT
    out.mkdir(parents=True)
    child = out / "CT-2026-077-1.xlsm"
    child.write_bytes(b"child")
    (out / "CT-2026-077.xlsm").write_bytes(b"parent")
    assert find_client_book(tmp_path, "2026-077-1") == child


def test_prepare_leaves_parent_book_and_template(tmp_path: Path) -> None:
    out = tmp_path / CLIENT_OUT
    out.mkdir(parents=True)
    template = out / "CT-202X-XXX.xlsm"
    template.write_bytes(b"blank")
    parent = out / "CT-2026-077.xlsm"
    parent.write_bytes(b"live")
    dest = prepare_client_book(tmp_path, "2026-077-1")
    assert dest == parent
    assert template.is_file()
    assert not (out / "CT-2026-077-1.xlsm").exists()


def test_prepare_renames_template_book(tmp_path: Path) -> None:
    out = tmp_path / CLIENT_OUT
    out.mkdir(parents=True)
    template = out / "CT-202X-XXX.xlsm"
    template.write_bytes(b"book")
    dest = prepare_client_book(tmp_path, "2026-Tanzim")
    assert dest.name == "CT-2026-Tanzim.xlsm"
    assert dest.is_file()
    assert not template.exists()


def test_file_client_copies_tab_and_bumps_number(tmp_path: Path) -> None:
    path = tmp_path / "CT-2026-Tanzim.xlsx"
    empty_client_workbook().save(path)
    lines = [
        LogLine("2026-Tanzim-1-1", "A", "Drawing-1", "APPROVAL"),
        LogLine("2026-Tanzim-1-2", "A", "Drawing-2", "APPROVAL"),
    ]
    result = file_client_transmittal(path, "2026-Tanzim", lines, issued=date(2026, 9, 1))
    assert result.cover_id == "CT-2026-Tanzim-1"
    assert result.sheet_name == "1"
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "TRANSMITTAL"
    assert "1" in wb.sheetnames
    filed = wb["1"]
    assert filed["A12"].value == "2026-Tanzim"
    assert filed["A15"].value == "2026-Tanzim-1-1"
    assert filed["C15"].value == "Drawing-1"
    assert filed["D15"].value == "APPROVAL"
    working = wb["TRANSMITTAL"]
    assert working["A2"].value == 2
    assert working["A15"].value is None
    wb.close()
    inspect = inspect_client_book(path, "2026-Tanzim")
    assert inspect.next_number == 2
    assert inspect.filed_tabs == ("1",)
    second = file_client_transmittal(
        path,
        "2026-Tanzim",
        [LogLine("2026-Tanzim-1-2", "0", "Drawing-2", "CONSTRUCTION")],
        issued=date(2026, 9, 2),
    )
    assert second.cover_id == "CT-2026-Tanzim-2"
    wb = load_workbook(path)
    assert wb.sheetnames[:3] == ["TRANSMITTAL", "2", "1"]
    wb.close()


def test_fill_writes_working_sheet_without_archive(tmp_path: Path) -> None:
    path = tmp_path / "CT-2026-Tanzim.xlsx"
    empty_client_workbook().save(path)
    pep = tmp_path / "pep.xlsx"
    cover = PepCover(
        path=pep,
        from_address=DOC_CONTROL_FROM,
        to_line="client@example.com",
        cc_line="pm@eliteintegrityservices.com",
        project_description="Client | Loc: Site | Ref. Tag: TK1 | PO#: 1 | WO#/MOC#: N/A",
        client="Client",
        site="Site",
        tank_tag="TK1",
        po="1",
        wo="",
    )
    result = fill_client_transmittal(
        path,
        "2026-Tanzim",
        [LogLine("2026-Tanzim-1-1", "A", "Drawing-1", "APPROVAL")],
        issued=date(2026, 9, 1),
        expected_return=date(2026, 9, 15),
        cover=cover,
    )
    assert result.cover_id == "CT-2026-Tanzim-1"
    assert result.sheet_name == "TRANSMITTAL"
    wb = load_workbook(path)
    assert wb.sheetnames == ["TRANSMITTAL"]
    working = wb["TRANSMITTAL"]
    assert working["A2"].value == 1
    assert working["A15"].value == "2026-Tanzim-1-1"
    assert working["C4"].value == DOC_CONTROL_FROM
    assert working["C12"].value.date() == date(2026, 9, 1)
    assert working["D12"].value.date() == date(2026, 9, 15)
    wb.close()
    inspect = inspect_client_book(path, "2026-Tanzim")
    assert inspect.next_number == 1
    assert inspect.filed_tabs == ()


def test_refuses_wrong_job_book(tmp_path: Path) -> None:
    path = tmp_path / "CT-2026-075.xlsx"
    wb = empty_client_workbook()
    wb["TRANSMITTAL"]["A12"] = "2026-075"
    wb.save(path)
    wb.close()
    try:
        inspect_client_book(path, "2026-Tanzim")
        raise AssertionError("should refuse")
    except LogError as exc:
        assert "2026-075" in str(exc)


def test_file_child_job_keeps_parent_book_number(tmp_path: Path) -> None:
    path = tmp_path / "CT-2026-077.xlsx"
    wb = empty_client_workbook()
    wb["TRANSMITTAL"]["A12"] = "2026-077"
    wb["TRANSMITTAL"]["A2"] = 4
    wb.save(path)
    wb.close()
    result = file_client_transmittal(
        path,
        "2026-077-1",
        [LogLine("2026-077-1-1", "0", "Drawing-1", "APPROVAL")],
        issued=date(2026, 9, 30),
    )
    assert result.cover_id == "CT-2026-077-4"
    inspect = inspect_client_book(path, "2026-077-1")
    assert inspect.job_number == "2026-077"
    assert inspect.cover_id == "CT-2026-077-5"
    wb = load_workbook(path)
    assert wb["TRANSMITTAL"]["A12"].value == "2026-077"
    assert wb["4"]["A12"].value == "2026-077"
    wb.close()


def test_file_real_template_copy_if_sandbox_exists(tmp_path: Path) -> None:
    # Filing a real .xlsm goes through Excel COM; skip in pytest (close the CT book first).
    return


def test_file_writes_pep_cover_and_return_date(tmp_path: Path) -> None:
    path = tmp_path / "CT-2026-Tanzim.xlsx"
    empty_client_workbook().save(path)
    pep = tmp_path / "2026-075 Project Execution Plan Rev2.xlsx"
    cover = PepCover(
        path=pep,
        from_address=DOC_CONTROL_FROM,
        to_line="glen.philips@interpipeline.com; rick.jondreau@outlook.com",
        cc_line="breydon@eliteintegrityservices.com",
        project_description=(
            "IPL | Loc: Kerrobert SK, 02-34-033-22-W3 | Ref. Tag: TK9 | PO#: 4000072534 | WO#/MOC#: N/A"
        ),
        client="IPL",
        site="Kerrobert SK, 02-34-033-22-W3",
        tank_tag="TK9",
        po="4000072534",
        wo="",
    )
    file_client_transmittal(
        path,
        "2026-Tanzim",
        [LogLine("2026-Tanzim-1-1", "A", "Drawing-1", "APPROVAL")],
        issued=date(2026, 9, 1),
        expected_return=date(2026, 9, 15),
        cover=cover,
    )
    wb = load_workbook(path)
    filed = wb["1"]
    assert filed["C4"].value == DOC_CONTROL_FROM
    assert filed["C6"].value == cover.to_line
    assert filed["C8"].value == cover.cc_line
    assert filed["A10"].value == cover.project_description
    assert filed["C12"].value.date() == date(2026, 9, 1)
    assert filed["D12"].value.date() == date(2026, 9, 15)
    working = wb["TRANSMITTAL"]
    assert working["C12"].value is None
    assert working["D12"].value is None
    wb.close()


def test_parse_dates() -> None:
    assert parse_issued_date("2026-09-01") == date(2026, 9, 1)
    assert parse_issued_date("") == date.today()
    assert parse_issued_date("N/A") == date.today()
    assert parse_issued_date("n/a") == date.today()
    assert parse_expected_return("") == "N/A"
    assert parse_expected_return("As soon as able") == "As soon as able"
    assert parse_expected_return("2026-09-15") == date(2026, 9, 15)
    assert cover_return_request_stamp("") == ""
    assert cover_return_request_stamp("N/A") == ""
    assert cover_return_request_stamp("As soon as able") == ""
    assert cover_return_request_stamp("2026-09-15") == "2026-09-15"
    assert cover_return_request_stamp("09/15/2026") == "2026-09-15"
    assert cover_date_stamp("") == ""
    assert cover_date_stamp("N/A") == ""
    assert cover_submission_date_stamp("") == ""
    assert cover_submission_date_stamp("2026-09-11") == "2026-09-11"
    assert cover_submission_date_stamp("09/11/2026") == "2026-09-11"
    assert cover_issued_default(today=date(2026, 9, 16)) == "2026-09-16"
    assert expected_return_from_issued("2026-09-16", 0) == "2026-09-16"
    assert expected_return_from_issued("2026-09-16", 1) == "2026-09-17"
    assert expected_return_from_issued("2026-09-16", 7) == "2026-09-23"
    assert expected_return_from_issued("2026-09-16", 14) == "2026-09-30"
    # Calendar days, not business days: Friday + 1 is Saturday.
    assert expected_return_from_issued("2026-09-18", 1) == "2026-09-19"
    assert expected_return_from_issued("N/A", 0, today=date(2026, 9, 16)) == "2026-09-16"
    assert expected_return_from_issued("", 1, today=date(2026, 9, 16)) == "2026-09-17"


def test_file_shop_uses_st_cells(tmp_path: Path) -> None:
    path = tmp_path / "ST-2026-Tanzim.xlsx"
    empty_shop_workbook().save(path)
    pep = tmp_path / "pep.xlsx"
    cover = PepCover(
        path=pep,
        from_address=DOC_CONTROL_FROM,
        to_line="client@example.com",
        cc_line="pm@eliteintegrityservices.com",
        project_description="Client | Loc: Site",
        client="Shop Co",
        site="Yard",
        tank_tag="",
        po="",
        wo="",
    )
    result = file_client_transmittal(
        path,
        "2026-Tanzim",
        [LogLine("EIS-1", "4", "SMAW", "CONSTRUCTION")],
        issued=date(2026, 9, 10),
        cover=cover,
        kind=SHOP,
    )
    assert result.cover_id == "ST-2026-Tanzim-1"
    wb = load_workbook(path)
    filed = wb["1"]
    assert filed["I4"].value == "2026-Tanzim"
    assert filed["C2"].value == 1
    assert filed["A6"].value == DOC_CONTROL_FROM
    assert filed["C6"].value is None
    assert filed["J4"].value == "Shop Co"
    assert filed["J6"].value == "Yard"
    assert filed["A9"].value == "EIS-1"
    assert filed["G9"].value == 4
    assert filed["I9"].value == "SMAW"
    assert filed["J9"].value == "CONSTRUCTION"
    working = wb["TRANSMITTAL"]
    assert working["C2"].value == 2
    assert working["A9"].value is None
    wb.close()
    inspect = inspect_book(path, "2026-Tanzim", SHOP)
    assert inspect.next_number == 2
    assert inspect.cover_id == "ST-2026-Tanzim-2"


def test_file_field_uses_ft_cells(tmp_path: Path) -> None:
    path = tmp_path / "FT-2026-Tanzim.xlsx"
    empty_field_workbook().save(path)
    result = file_client_transmittal(
        path,
        "2026-Tanzim",
        [LogLine("2026-Tanzim-1-1", "0", "Drawing-1", "INFORMATION")],
        issued=date(2026, 9, 10),
        kind=FIELD,
    )
    assert result.cover_id == "FT-2026-Tanzim-1"
    wb = load_workbook(path)
    filed = wb["1"]
    assert filed["I6"].value == "2026-Tanzim"
    assert filed["I4"].value is None
    assert filed["A8"].value == DOC_CONTROL_FROM
    assert filed["A10"].value == "2026-Tanzim-1-1"
    assert filed["I10"].value == "Drawing-1"
    assert filed["J10"].value == "INFORMATION"
    working = wb["TRANSMITTAL"]
    assert working["C2"].value == 2
    wb.close()


def test_read_book_cover_uses_existing_letter(tmp_path: Path) -> None:
    path = tmp_path / "CT-2026-Tanzim.xlsx"
    wb = empty_client_workbook()
    ws = wb["TRANSMITTAL"]
    ws["C6"] = "Zach.Hilsendager@cnrl.com"
    ws["C8"] = "tanzim.nasir@eliteintegrityservices.com; sarah.chan@eliteintegrityservices.com"
    ws["A10"] = "Canadian Natural Resources Ltd. | Loc: Alberta | Ref. Tag: To be confirmed"
    wb.save(path)
    wb.close()
    cover = read_book_cover(path)
    assert cover.to_line == "Zach.Hilsendager@cnrl.com"
    assert "sarah.chan@eliteintegrityservices.com" in cover.cc_line
    assert cover.project_description.startswith("Canadian Natural Resources Ltd.")


def test_pick_cover_prefers_letter_over_pep() -> None:
    pep = PepCover(
        path=Path("pep.xlsx"),
        from_address=DOC_CONTROL_FROM,
        to_line="pep@example.com",
        cc_line="cc@example.com",
        project_description="IPL | Loc: Site | Ref. Tag: TK9 | PO#: 1 | WO#/MOC#: N/A",
        client="IPL",
        site="Site",
        tank_tag="TK9",
        po="1",
        wo="",
    )
    book = BookCover(
        to_line="Zach.Hilsendager@cnrl.com",
        cc_line="sarah.chan@eliteintegrityservices.com",
        project_description="CNRL letter line",
    )
    chosen = pick_cover_fields(
        book=book,
        pack_to="tanzim.nasir@eliteintegrityservices.com",
        pack_cc="",
        pack_project="saved json",
        pep=pep,
    )
    assert chosen.to_line == "Zach.Hilsendager@cnrl.com"
    assert chosen.cc_line == "sarah.chan@eliteintegrityservices.com"
    assert chosen.project_description == "CNRL letter line"


def test_pick_cover_falls_back_to_pep_when_letter_blank() -> None:
    pep = PepCover(
        path=Path("pep.xlsx"),
        from_address=DOC_CONTROL_FROM,
        to_line="pep@example.com",
        cc_line="",
        project_description="From PEP",
        client="IPL",
        site="",
        tank_tag="",
        po="",
        wo="",
    )
    chosen = pick_cover_fields(book=BookCover(), pep=pep)
    assert chosen.to_line == "pep@example.com"
    assert chosen.project_description == "From PEP"


def test_field_cover_uses_settings_when_letter_and_pack_are_blank() -> None:
    chosen = pick_cover_fields(
        book=BookCover(),
        kind=FIELD,
        field_to="foreman@eliteintegrityservices.com",
        field_cc="super@eliteintegrityservices.com",
    )
    assert chosen.to_line == "foreman@eliteintegrityservices.com"
    assert chosen.cc_line == "super@eliteintegrityservices.com"
    letter = pick_cover_fields(
        book=BookCover(to_line="already@eliteintegrityservices.com", cc_line=""),
        kind=FIELD,
        field_to="foreman@eliteintegrityservices.com",
        field_cc="super@eliteintegrityservices.com",
    )
    assert letter.to_line == "already@eliteintegrityservices.com"
    assert letter.cc_line == "super@eliteintegrityservices.com"


def test_shop_cover_uses_settings_when_letter_and_pack_are_blank() -> None:
    chosen = pick_cover_fields(
        book=BookCover(),
        kind=SHOP,
        shop_to="shop@eliteintegrityservices.com",
        shop_cc="buyer@eliteintegrityservices.com",
    )
    assert chosen.to_line == "shop@eliteintegrityservices.com"
    assert chosen.cc_line == "buyer@eliteintegrityservices.com"
    letter = pick_cover_fields(
        book=BookCover(to_line="already@eliteintegrityservices.com", cc_line=""),
        kind=SHOP,
        shop_to="shop@eliteintegrityservices.com",
        shop_cc="buyer@eliteintegrityservices.com",
    )
    assert letter.to_line == "already@eliteintegrityservices.com"
    assert letter.cc_line == "buyer@eliteintegrityservices.com"
