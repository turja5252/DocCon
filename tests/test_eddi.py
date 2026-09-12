# Created by Tanzim Nasir
# Copyright (c) 2026 Tanzim Nasir.
# Built for Elite Integrity Services.
# Unauthorized use by other companies is prohibited.
from datetime import date
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.pagebreak import Break

from doccon.client_log import LogError
from doccon.eddi import (
    CATEGORY_HEADER_MERGES,
    find_live_eddi_book,
    form_group_items,
    grouped_eddi_items,
    is_dated_eddi_name,
    parse_form_groups,
    snapshot_eddi,
)
from doccon.register import DrawingRow

ORANGE = PatternFill(fill_type="solid", fgColor="F79646")

FORM_GROUPS = (
    "FABRICATION DRAWINGS",
    "CONSTRUCTION / INSTALLATION DRAWINGS",
    "SUPPLIER DRAWINGS",
    "ENGINEERING",
    "QC - WO, ITP, NDE RECORDS, PROCEDURES",
    "QC - WELDING",
    "QC - SUPPLIER",
    "DOCUMENT CONTROL",
    "PROJECT MANAGEMENT",
)


def _drawing(**kwargs) -> DrawingRow:
    values = {
        "key": "P2024-1",
        "summary": "2026-Tanzim-1-1 Drawing-1",
        "drawing_id": "2026-Tanzim-1-1",
        "title": "Drawing-1",
        "status": "To Do",
        "job_number": "2026-Tanzim",
        "outgoing_rev": "B",
        "purpose": "Approval",
        "parent_summary": "Drawing Package",
        "eddi_status": "1 - Fabrication Drawings - EDDI",
        "submission_date": "2026-09-10",
    }
    values.update(kwargs)
    return DrawingRow(**values)


def _project_fixture(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Project"
    ws["A1"] = "ELITE DRAWING & DOCUMENT INDEX"
    ws.merge_cells("A1:M1")
    ws["A2"] = "EDDI-2026-XXXX"
    ws.merge_cells("A2:M2")
    ws["A3"] = "DOCUMENT NUMBER"
    ws["B3"] = "CLIENT DOCUMENT NUMBER"
    ws["C3"] = "DOCUMENT DESCRIPTION"
    ws["D3"] = "OUTGOING"
    ws["G3"] = "INCOMING"
    ws["J3"] = "SHOP IFC"
    ws["L3"] = "FIELD IFC"
    ws["D4"] = "REV"
    ws["E4"] = "SUBMITTED TO CLIENT FOR "
    ws["F4"] = "SUBMISSION DATE"
    ws["G4"] = "REV"
    ws["H4"] = "CLIENT APPROVAL STATUS"
    ws["I4"] = "RETURNED DATE"
    ws["J4"] = "REV"
    ws["K4"] = "DATE"
    ws["L4"] = "REV"
    ws["M4"] = "DATE"
    row = 5
    for number, title in enumerate(FORM_GROUPS, start=1):
        for col in range(1, 14):
            cell = ws.cell(row, col)
            cell.fill = ORANGE
            cell.alignment = Alignment(horizontal="center")
        ws.cell(row, 1, f"{number}. {title}")
        row += 1
        ws.cell(row, 1, "Column1")
        for col in range(2, 14):
            ws.cell(row, col, f"Column{col}")
        ws.row_dimensions[row].hidden = True
        row += 1
        if number == 6:
            ws.cell(row, 1, "WS-2024-XXX")
            ws.cell(row, 3, "Elite Weld Procedure Summary")
            ws.cell(row, 4, "A")
        row += 1
        row += 1
    last = row - 1
    ws.print_area = f"A1:M{last}"
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.row_breaks.append(Break(id=24, man=True, max=13))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    return path


def _merged(ws, ref: str) -> bool:
    min_col, min_row, max_col, max_row = range_boundaries(ref)
    target = (min_col, min_row, max_col, max_row)
    return any(
        (m.min_col, m.min_row, m.max_col, m.max_row) == target for m in ws.merged_cells.ranges
    )


def _hidden(ws, row: int) -> bool:
    return bool(ws.row_dimensions[row].hidden)


def test_grouped_eddi_skips_generic_and_empty_blocks() -> None:
    generic = _drawing(key="P2024-0", eddi_status="0 - Generic Task", drawing_id="GEN")
    weld = _drawing(
        key="P2024-3",
        drawing_id="EIS-1",
        title="WPS",
        eddi_status="6 - QC - Welding - EDDI",
    )
    fab = _drawing()
    groups = grouped_eddi_items([generic, weld, fab])
    assert [title for title, _items in groups] == ["1 - Fabrication Drawings", "6 - QC - Welding"]
    assert [row.drawing_id for _title, items in groups for row in items] == ["2026-Tanzim-1-1", "EIS-1"]


def test_form_groups_skip_generic_and_ungrouped() -> None:
    buckets = form_group_items(
        [
            _drawing(),
            _drawing(key="P2024-0", eddi_status="0 - Generic Task", drawing_id="GEN"),
            _drawing(key="P2024-9", eddi_status="", drawing_id="LOOSE"),
            _drawing(key="P2024-3", drawing_id="EIS-1", eddi_status="6 - QC - Welding - EDDI"),
        ]
    )
    assert [row.drawing_id for row in buckets[1]] == ["2026-Tanzim-1-1"]
    assert [row.drawing_id for row in buckets[6]] == ["EIS-1"]
    assert buckets[2] == []
    assert all(not buckets[n] or all(eddi_group_ok(row) for row in buckets[n]) for n in buckets)


def eddi_group_ok(row: DrawingRow) -> bool:
    return row.drawing_id not in {"GEN", "LOOSE"}


def test_find_live_book_ignores_dated_leftover(tmp_path: Path) -> None:
    folder = tmp_path / "2026-Tanzim" / "3.0 Doc Con"
    leftover = folder / "EDDI-2026-Tanzim-2026-09-08.xlsx"
    leftover.parent.mkdir(parents=True)
    leftover.write_bytes(b"old-custom-list")
    live = _project_fixture(folder / "EDDI-2026-Tanzim.xlsx")
    found = find_live_eddi_book(tmp_path / "2026-Tanzim", "2026-Tanzim")
    assert found == live
    assert is_dated_eddi_name(leftover.name, "2026-Tanzim")
    assert not is_dated_eddi_name(live.name, "2026-Tanzim")


def test_find_live_book_missing_raises(tmp_path: Path) -> None:
    folder = tmp_path / "2026-Tanzim" / "3.0 Doc Con"
    folder.mkdir(parents=True)
    (folder / "EDDI-2026-Tanzim-2026-09-08.xlsx").write_bytes(b"leftover")
    with pytest.raises(LogError, match="EDDI-2026-Tanzim.xlsm"):
        find_live_eddi_book(tmp_path / "2026-Tanzim", "2026-Tanzim")


def test_snapshot_fills_project_form_not_custom_list(tmp_path: Path) -> None:
    job_folder = tmp_path / "2026-Tanzim"
    live = _project_fixture(job_folder / "3.0 Doc Con" / "EDDI-2026-Tanzim.xlsx")
    before = live.read_bytes()
    drawings = [
        _drawing(
            client_document_number="CNRL-1",
            due_date="2099-01-01",
            return_request_date="2099-02-02",
            incoming_rev="A",
            approval="Approved",
            return_date="2026-09-30",
            shop_ifc_rev="0",
            shop_ifc_date="2026-09-01",
            field_ifc_rev="0",
            field_ifc_date="2026-09-02",
        ),
        _drawing(
            key="P2024-3",
            drawing_id="EIS-1",
            title="Welding Procedure Specification",
            outgoing_rev="4",
            purpose="",
            submission_date="",
            eddi_status="6 - QC - Welding - EDDI",
        ),
        _drawing(key="P2024-9", eddi_status="0 - Generic Task", drawing_id="SKIP"),
        _drawing(key="P2024-8", eddi_status="", drawing_id="LOOSE"),
        _drawing(
            key="P2024-91",
            drawing_id="PM-1",
            title="Schedule",
            eddi_status="9 - Project Management - EDDI",
        ),
        _drawing(
            key="P2024-92",
            drawing_id="PM-2",
            title="Minutes",
            eddi_status="9 - Project Management - EDDI",
        ),
        _drawing(
            key="P2024-93",
            drawing_id="PM-3",
            title="Agenda",
            eddi_status="9 - Project Management - EDDI",
        ),
    ]
    snap = snapshot_eddi(job_folder, "2026-Tanzim", drawings, date(2026, 9, 10), print_pdf=False)
    assert snap.book.name == "EDDI-2026-Tanzim-2026-09-10.xlsx"
    assert snap.book != live
    assert live.read_bytes() == before
    assert snap.updated_rows == 5
    assert snap.note.startswith("EDDI: EDDI-2026-Tanzim-2026-09-10.pdf")

    wb = load_workbook(snap.book)
    assert wb.sheetnames[0] == "Project"
    assert "EDDI" not in wb.sheetnames
    ws = wb["Project"]
    assert ws["A1"].value == "ELITE DRAWING & DOCUMENT INDEX"
    assert str(ws["A2"].value) == "EDDI-2026-Tanzim"
    groups = parse_form_groups(ws)
    assert [group.number for group in groups] == list(range(1, 10))
    assert ws["A3"].value == "DOCUMENT NUMBER"
    values_a = [ws.cell(row, 1).value for row in range(1, ws.max_row + 1)]
    assert "1. FABRICATION DRAWINGS" in values_a
    assert "6. QC - WELDING" in values_a
    assert "9. PROJECT MANAGEMENT" in values_a
    assert "2026-Tanzim-1-1" in values_a
    assert "EIS-1" in values_a
    assert "PM-1" in values_a
    assert "PM-2" in values_a
    assert "PM-3" in values_a
    assert "SKIP" not in values_a
    assert "LOOSE" not in values_a
    assert "WS-2024-XXX" not in values_a
    item = next(row for row in range(1, ws.max_row + 1) if ws.cell(row, 1).value == "2026-Tanzim-1-1")
    assert ws.cell(item, 2).value == "CNRL-1"
    assert ws.cell(item, 3).value == "Drawing-1"
    assert ws.cell(item, 4).value == "B"
    assert ws.cell(item, 5).value == "Approval"
    assert str(ws.cell(item, 6).value)[:10] == "2026-09-10"
    assert ws.cell(item, 7).value == "A"
    assert ws.cell(item, 8).value == "Approved"
    assert str(ws.cell(item, 9).value)[:10] == "2026-09-30"
    assert ws.cell(item, 10).value == "0"
    assert str(ws.cell(item, 11).value)[:10] == "2026-09-01"
    assert ws.cell(item, 12).value == "0"
    assert str(ws.cell(item, 13).value)[:10] == "2026-09-02"
    blob = " ".join(str(ws.cell(row, col).value or "") for row in range(1, ws.max_row + 1) for col in range(1, 14))
    assert "To Do" not in blob
    assert "2099-01-01" not in blob
    assert "2099-02-02" not in blob
    assert "1 - Fabrication Drawings - EDDI" not in blob
    assert "INCLUDE IN DATA BOOK" not in blob
    junk_rows = [row for row in range(1, ws.max_row + 1) if str(ws.cell(row, 1).value or "").startswith("Column1")]
    assert junk_rows
    assert all(ws.row_dimensions[row].hidden for row in junk_rows)
    fab = next(group for group in groups if group.number == 1)
    assert not ws.row_dimensions[fab.first_item].hidden
    assert ws.row_dimensions[fab.first_item + 1].hidden
    pm = next(group for group in groups if group.number == 9)
    assert pm.slot_count == 3
    assert [ws.cell(row, 1).value for row in range(pm.first_item, pm.last_item + 1)] == ["PM-1", "PM-2", "PM-3"]
    by_num = {group.number: group for group in groups}
    for number in (1, 6, 9):
        header = by_num[number].header_row
        assert not _hidden(ws, header)
        assert _merged(ws, f"A{header}:M{header}")
        align = ws.cell(header, 1).alignment
        assert align.horizontal == "center"
        assert align.vertical == "center"
        fill = ws.cell(header, 1).fill
        assert fill.fill_type == "solid"
        assert str(fill.fgColor.rgb).endswith("F79646")
        assert str(ws.cell(header, 1).value).startswith(f"{number}.")
    for number in (2, 3, 4, 5, 7, 8):
        header = by_num[number].header_row
        assert _hidden(ws, header)
        if by_num[number].junk_row is not None:
            assert _hidden(ws, by_num[number].junk_row)
        if by_num[number].slot_count:
            assert all(
                _hidden(ws, row) for row in range(by_num[number].first_item, by_num[number].last_item + 1)
            )
    for ref in CATEGORY_HEADER_MERGES:
        assert _merged(ws, ref)
        min_col, min_row, _max_col, _max_row = range_boundaries(ref)
        cat = ws.cell(min_row, min_col).alignment
        assert cat.horizontal == "center"
        assert cat.vertical == "center"
    assert 24 in [int(brk.id) for brk in ws.row_breaks.brk]
    wb.close()
    xml = ZipFile(snap.book).read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert 't="inlineStr"></c>' not in xml


def test_snapshot_leaves_live_template_untouched(tmp_path: Path) -> None:
    job_folder = tmp_path / "2026-Tanzim"
    live = _project_fixture(job_folder / "3.0 Doc Con" / "EDDI-2026-Tanzim.xlsx")
    leftover = job_folder / "3.0 Doc Con" / "EDDI-202X-XXX.xlsm"
    leftover.write_bytes(b"template")
    snapshot_eddi(job_folder, "2026-Tanzim", [_drawing()], date(2026, 9, 10), print_pdf=False)
    assert leftover.read_bytes() == b"template"
    wb = load_workbook(live)
    assert wb["Project"]["A2"].value == "EDDI-2026-XXXX"
    wb.close()


def test_snapshot_hides_empty_groups_and_drops_stranded_break(tmp_path: Path) -> None:
    job_folder = tmp_path / "2026-Tanzim"
    _project_fixture(job_folder / "3.0 Doc Con" / "EDDI-2026-Tanzim.xlsx")
    drawings = [
        _drawing(
            key="P2024-3",
            drawing_id="EIS-1",
            title="WPS",
            eddi_status="6 - QC - Welding - EDDI",
        ),
        _drawing(
            key="P2024-91",
            drawing_id="PM-1",
            title="Schedule",
            eddi_status="9 - Project Management - EDDI",
        ),
    ]
    snap = snapshot_eddi(job_folder, "2026-Tanzim", drawings, date(2026, 9, 11), print_pdf=False)
    wb = load_workbook(snap.book)
    ws = wb["Project"]
    groups = parse_form_groups(ws)
    by_num = {group.number: group for group in groups}
    assert not _hidden(ws, by_num[6].header_row)
    assert not _hidden(ws, by_num[9].header_row)
    assert _merged(ws, f"A{by_num[6].header_row}:M{by_num[6].header_row}")
    assert _merged(ws, f"A{by_num[9].header_row}:M{by_num[9].header_row}")
    assert ws.cell(by_num[6].header_row, 1).value == "6. QC - WELDING"
    assert ws.cell(by_num[9].header_row, 1).value == "9. PROJECT MANAGEMENT"
    for number in (1, 2, 3, 4, 5, 7, 8):
        header = by_num[number].header_row
        assert _hidden(ws, header)
        assert not _merged(ws, f"A{header}:M{header}")
    assert [int(brk.id) for brk in ws.row_breaks.brk] == []
    wb.close()


def test_snapshot_includes_matched_pdf_only(tmp_path: Path) -> None:
    from doccon.eddi import eddi_print_drawings
    from doccon.match import MatchedRow, PdfHit

    job_folder = tmp_path / "2026-Tanzim"
    _project_fixture(job_folder / "3.0 Doc Con" / "EDDI-2026-Tanzim.xlsx")
    pdf = tmp_path / "2026-Tanzim-1-1 REV A.pdf"
    pdf.write_bytes(b"%PDF")
    matched = MatchedRow(
        drawing=_drawing(),
        pdf=PdfHit(path=pdf, drawing_id="2026-Tanzim-1-1", rev="A"),
        confidence="High",
    )
    missing = MatchedRow(
        drawing=_drawing(
            key="P2024-2",
            drawing_id="2026-Tanzim-1-2",
            title="Placeholder",
            eddi_status="1 - Fabrication Drawings - EDDI",
        ),
        pdf=None,
        confidence="Missing",
    )
    drawings = eddi_print_drawings([matched, missing])
    assert [row.drawing_id for row in drawings] == ["2026-Tanzim-1-1"]
    snap = snapshot_eddi(job_folder, "2026-Tanzim", drawings, date(2026, 9, 11), print_pdf=False)
    assert snap.updated_rows == 1
    wb = load_workbook(snap.book)
    values = [ws_cell for ws_cell in (wb["Project"].cell(row, 1).value for row in range(1, wb["Project"].max_row + 1))]
    wb.close()
    assert "2026-Tanzim-1-1" in values
    assert "2026-Tanzim-1-2" not in values
