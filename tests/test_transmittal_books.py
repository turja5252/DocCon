# Created by Tanzim Nasir
# Copyright (c) 2026 Tanzim Nasir.
# Built for Elite Integrity Services.
# Unauthorized use by other companies is prohibited.
from pathlib import Path

from openpyxl import Workbook, load_workbook

from doccon.kinds import BOOK_DIR, CLIENT, FIELD, SHOP
from doccon.transmittal_books import adopt_transmittal_books, is_template_book_name


def test_template_name_detects_placeholders() -> None:
    assert is_template_book_name("CT-202X-XXX.xlsm", "CT")
    assert is_template_book_name("ST-20XX-XXX.xlsm", "ST")
    assert is_template_book_name("FT-20XX-XXX.xlsm", "FT")
    assert not is_template_book_name("CT-2026-Tanzim.xlsm", "CT")
    assert not is_template_book_name("CT-2026-075.xlsm", "CT")
    assert not is_template_book_name("ST-2026-Tanzim.xlsm", "ST")


def test_adopt_renames_all_three_and_stamps_job(tmp_path: Path) -> None:
    job = tmp_path / "2026-Tanzim"
    (job / BOOK_DIR[CLIENT]).mkdir(parents=True)
    (job / BOOK_DIR[SHOP]).mkdir(parents=True)
    (job / BOOK_DIR[FIELD]).mkdir(parents=True)

    def blank(path: Path, cell: str) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "TRANSMITTAL"
        ws[cell] = "202X-XXX"
        wb.save(path)
        wb.close()

    blank(job / BOOK_DIR[CLIENT] / "CT-202X-XXX.xlsm", "A12")
    blank(job / BOOK_DIR[SHOP] / "ST-20XX-XXX.xlsm", "I4")
    blank(job / BOOK_DIR[FIELD] / "FT-20XX-XXX.xlsm", "I6")

    adopted = adopt_transmittal_books(job, "2026-Tanzim")
    names = sorted(item.path.name for item in adopted)
    assert names == ["CT-2026-Tanzim.xlsm", "FT-2026-Tanzim.xlsm", "ST-2026-Tanzim.xlsm"]
    assert not (job / BOOK_DIR[CLIENT] / "CT-202X-XXX.xlsm").exists()

    ct = load_workbook(job / BOOK_DIR[CLIENT] / "CT-2026-Tanzim.xlsm")
    assert ct["TRANSMITTAL"]["A12"].value == "2026-Tanzim"
    ct.close()
    st = load_workbook(job / BOOK_DIR[SHOP] / "ST-2026-Tanzim.xlsm")
    assert st["TRANSMITTAL"]["I4"].value == "2026-Tanzim"
    st.close()
    ft = load_workbook(job / BOOK_DIR[FIELD] / "FT-2026-Tanzim.xlsm")
    assert ft["TRANSMITTAL"]["I6"].value == "2026-Tanzim"
    ft.close()

    again = adopt_transmittal_books(job, "2026-Tanzim")
    assert again == []


def test_adopt_leaves_live_named_book_alone(tmp_path: Path) -> None:
    job = tmp_path / "2026-075"
    folder = job / BOOK_DIR[CLIENT]
    folder.mkdir(parents=True)
    live = folder / "CT-2026-075.xlsm"
    live.write_bytes(b"live")
    adopted = adopt_transmittal_books(job, "2026-075")
    assert adopted == []
    assert live.is_file()
    assert live.read_bytes() == b"live"
